"""Разбор списка деталей из выгрузки заказчика.

Формат подсмотрен в реальном файле «Расчёт заказов» студии фасадов:

* шапка с датой и заказчиком;
* строка заголовков ``№ | Длина | Ширина | Кол-во | М2 | Мм | Фрезеровка | Цвет``;
* **два блока позиций рядом на одном листе** — второй начинается правее, со
  своим таким же заголовком. Поэтому ищем не одну таблицу, а все;
* внизу строка итогов, которую в детали брать нельзя.

Разбор терпимый к мелочам: колонки могут идти в другом порядке, называться
чуть иначе («Кол-во» / «Кол» / «Шт»), а количество — отсутствовать.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from app.core.material_text import clean_color

SQ_MM_IN_SQ_M = Decimal("1000000")

# Заголовки колонок. Ключ — поле детали, значения — что может быть написано.
COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "width": ("длина", "длинна", "length"),
    "height": ("ширина", "width"),
    "quantity": ("колво", "кол", "количество", "шт", "штук", "qty"),
    "area": ("м2", "м²", "площадь", "кв", "area"),
    "thickness": ("мм", "толщина", "thickness"),
    "milling": ("фрезеровка", "фрезировка", "фрез", "milling"),
    "color": ("цвет", "декор", "color"),
}

# Строки, которые не являются позициями
TOTAL_MARKERS = ("итог", "всего", "размеры проверил", "заказ принял", "студия фасадов")

MAX_ROWS = 5000
MAX_COLUMNS = 60

# «RAL 9003», «RAL9003», «рал 9003» — заказчик пишет по-разному, а в заказ
# должно попасть одно каноничное написание.
RAL_PATTERN = re.compile(r"\b(?:RAL|РАЛ)\s*-?\s*(\d{3,4})\b", re.IGNORECASE)

# Толщина внутри текста: «плита 19мм», «16 мм», «19,5 мм».
# Взгляд назад обязателен: без него из «9003мм» вырезалось бы «003».
THICKNESS_IN_TEXT_PATTERN = re.compile(
    r"(?<![\d.,])(\d{1,3}(?:[.,]\d+)?)\s*мм",
    re.IGNORECASE,
)

# Разумные пределы толщины плиты: всё вне их — не толщина, а случайное число.
MIN_THICKNESS_MM = Decimal("1")
MAX_THICKNESS_MM = Decimal("100")


@dataclass
class ParsedPart:
    width: int
    height: int
    quantity: int
    thickness: str | None = None
    milling: str | None = None
    color: str | None = None
    # исходные значения из бланка — нужны менеджеру для сверки с бумагой,
    # поэтому чистка не имеет права их затирать
    color_source: str | None = None
    milling_source: str | None = None

    @property
    def area_sqm(self) -> Decimal:
        return (Decimal(self.width) * Decimal(self.height) * Decimal(self.quantity)) / SQ_MM_IN_SQ_M


@dataclass
class ParseResult:
    parts: list[ParsedPart] = field(default_factory=list)
    customer: str | None = None
    order_number: str | None = None
    skipped_rows: int = 0
    sheet_name: str | None = None

    @property
    def total_quantity(self) -> int:
        return sum(part.quantity for part in self.parts)

    @property
    def total_area_sqm(self) -> Decimal:
        return sum((part.area_sqm for part in self.parts), Decimal("0"))


class PartsImportError(Exception):
    """Файл не удалось разобрать — сообщение показывается менеджеру."""


def _normalize(value: object) -> str:
    """Приводит заголовок к виду, по которому его можно узнать."""
    text = str(value or "").strip().lower()
    return re.sub(r"[^a-zа-я0-9]", "", text)


def _clean(value: object) -> str:
    return str(value or "").strip()


def _to_int(value: object) -> int | None:
    """Числа приходят строками, дробями и с запятой — приводим бережно."""
    text = _clean(value).replace(",", ".").replace(" ", "").replace(" ", "")
    if not text:
        return None
    try:
        number = Decimal(text)
    except (InvalidOperation, ValueError):
        return None
    if number <= 0:
        return None
    return int(number.to_integral_value())


def _collapse_spaces(value: object) -> str:
    """Обрезает края и схлопывает двойные пробелы и переносы строк."""
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _format_thickness(number: Decimal) -> str:
    """«19.0» → «19», «19.50» → «19.5»: в заказ идёт число, а не запись из ячейки."""
    return format(number.normalize(), "f")


def _thickness_from_text(value: object) -> str | None:
    """Толщина из свободного текста: «Эмаль RAL 9003 плита 19мм» → «19».

    Годится и для колонки «Мм», и для названия цвета: в бланке заказчика
    колонка пустая, а толщина написана внутри цвета.
    """
    text = _collapse_spaces(value)
    if not text:
        return None

    candidates: list[str] = [match.group(1) for match in THICKNESS_IN_TEXT_PATTERN.finditer(text)]
    if not candidates:
        # в колонке «Мм» пишут голое число, без единиц измерения
        bare = text.replace(",", ".")
        if re.fullmatch(r"\d{1,3}(?:\.\d+)?", bare):
            candidates = [bare]
    for candidate in candidates:
        try:
            number = Decimal(candidate.replace(",", "."))
        except InvalidOperation:
            continue
        if MIN_THICKNESS_MM <= number <= MAX_THICKNESS_MM:
            return _format_thickness(number)
    return None


def _read_thickness(column_value: object, color_value: object) -> str | None:
    """Колонка «Мм» главнее: если её заполнили, текст цвета не спрашиваем."""
    from_column = _thickness_from_text(column_value)
    if from_column is not None:
        return from_column
    return _thickness_from_text(color_value)


def _clean_color(value: object) -> str | None:
    """Цвет в том виде, как его решил хранить заказчик.

    Правило общее для всего проекта — см. ``app.core.material_text``: тем же
    текстом приходит и загруженный файл, и строка, набранная руками.
    """
    return clean_color(value)


def _clean_milling(value: object) -> str | None:
    """Одну и ту же фрезеровку пишут «фрез», «Фрез», «фрез  » — сводим к одному.

    Только написание: значения заказчика не переводим и не подменяем своими,
    иначе технолог не узнает в заказе то, что написал на бланке.
    """
    text = _collapse_spaces(value)
    if not text:
        return None
    return text[0].lower() + text[1:]


def _is_total_row(cells: list[object]) -> bool:
    joined = " ".join(_clean(cell).lower() for cell in cells)
    return any(marker in joined for marker in TOTAL_MARKERS)


def _infer_quantity_column(block: dict[str, int]) -> None:
    """Достраивает колонку количества, если у неё нет заголовка.

    В бланке заказчика количество стоит между «Ширина» и «М2», а ячейка
    заголовка над ним пустая. Без этого весь заказ читался бы как «по одной
    штуке» — 37 деталей вместо 42.
    """
    if "quantity" in block:
        return
    height_index = block.get("height")
    area_index = block.get("area")
    if height_index is None or area_index is None:
        return
    candidate = height_index + 1
    if candidate < area_index and candidate not in block.values():
        block["quantity"] = candidate


def _find_header_blocks(rows: list[list[object]]) -> tuple[int, list[dict[str, int]]]:
    """Ищет строку заголовков и все блоки колонок в ней.

    Возвращает номер строки и список блоков: у каждого — своё отображение
    «поле детали → индекс колонки». Блоков может быть несколько, потому что в
    бланке заказчика позиции идут двумя таблицами рядом.
    """
    for row_index, row in enumerate(rows[:60]):
        found: list[tuple[int, str]] = []
        for column_index, cell in enumerate(row):
            key = _normalize(cell)
            if not key:
                continue
            for field_name, aliases in COLUMN_ALIASES.items():
                if key in aliases:
                    found.append((column_index, field_name))
                    break

        if not any(name == "width" for _, name in found) or not any(
            name == "height" for _, name in found
        ):
            continue

        # Блок закрывается там, где колонка повторяется: в бланке заказчика
        # второй набор «Длина/Ширина/...» означает вторую таблицу на том же
        # листе. Проверять именно повтор, а не «встретилась Длина», обязательно:
        # порядок колонок в чужих выгрузках бывает любой.
        blocks: list[dict[str, int]] = []
        current: dict[str, int] = {}
        for column_index, field_name in found:
            if field_name in current:
                blocks.append(current)
                current = {}
            current[field_name] = column_index
        if current:
            blocks.append(current)

        usable = [block for block in blocks if "width" in block and "height" in block]
        for block in usable:
            _infer_quantity_column(block)
        if usable:
            return row_index, usable

    raise PartsImportError(
        "Не нашёл строку с заголовками «Длина» и «Ширина». "
        "Проверьте, что в файле есть таблица с размерами."
    )


def _read_meta(rows: list[list[object]], header_row: int) -> tuple[str | None, str | None]:
    """Заказчик и номер бланка из шапки над таблицей."""
    customer: str | None = None
    order_number: str | None = None
    for row in rows[:header_row]:
        for index, cell in enumerate(row):
            key = _normalize(cell)
            neighbours = [_clean(item) for item in row[index + 1 : index + 4] if _clean(item)]
            if key.startswith("заказчик") and neighbours and customer is None:
                customer = neighbours[0]
            if key.startswith("бланкзаказа") and neighbours and order_number is None:
                order_number = neighbours[0]
    return customer, order_number


def _cell(row: list[object], block: dict[str, int], field_name: str) -> object:
    index = block.get(field_name)
    return row[index] if index is not None and index < len(row) else None


def _read_part(row: list[object], block: dict[str, int]) -> ParsedPart | None:
    """Одна деталь из строки. ``None`` — строка пустая, это не ошибка."""
    width = _to_int(_cell(row, block, "width"))
    height = _to_int(_cell(row, block, "height"))
    if width is None or height is None:
        # пустых строк в бланке десятки — заказчик печатает его с запасом
        return None

    thickness_cell = _cell(row, block, "thickness")
    milling_cell = _cell(row, block, "milling")
    color_cell = _cell(row, block, "color")

    return ParsedPart(
        width=width,
        height=height,
        quantity=_to_int(_cell(row, block, "quantity")) or 1,
        thickness=_read_thickness(thickness_cell, color_cell),
        milling=_clean_milling(milling_cell),
        color=_clean_color(color_cell),
        color_source=_clean(color_cell) or None,
        milling_source=_clean(milling_cell) or None,
    )


def _milling_key(value: str | None) -> str:
    """Сравнивает написания артикула, сохраняя значимое описание после номера."""
    text = _collapse_spaces(value).casefold()
    return re.sub(
        r"^(?:арт(?:икул)?|art(?:icle)?)[.\s№]*(?=\d)",
        "арт ",
        text,
    )


def _combine_identical_parts(parts: list[ParsedPart]) -> list[ParsedPart]:
    """Суммирует штуки одинаковых деталей по всему бланку, независимо от расстояния."""
    combined: list[ParsedPart] = []
    by_identity: dict[tuple[int, int, str | None, str, str], ParsedPart] = {}
    for part in parts:
        identity = (
            part.width,
            part.height,
            part.thickness,
            _milling_key(part.milling),
            _collapse_spaces(part.color).casefold(),
        )
        existing = by_identity.get(identity)
        if existing is None:
            by_identity[identity] = part
            combined.append(part)
        else:
            existing.quantity += part.quantity
    return combined


def parse_rows(rows: list[list[object]], sheet_name: str | None = None) -> ParseResult:
    """Разбирает уже прочитанные строки таблицы."""
    header_row, blocks = _find_header_blocks(rows)
    customer, order_number = _read_meta(rows, header_row)
    result = ParseResult(customer=customer, order_number=order_number, sheet_name=sheet_name)
    # The paper form prints material, thickness and milling once for a run of
    # details. Keep continuation state per block: some files put two separate
    # tables side by side on the same sheet.
    previous_by_block: list[ParsedPart | None] = [None] * len(blocks)

    for row in rows[header_row + 1 :]:
        if _is_total_row(row):
            # итоги и подписи — дальше позиций нет
            break
        for block_index, block in enumerate(blocks):
            part = _read_part(row, block)
            if part is not None:
                previous = previous_by_block[block_index]
                if previous is not None:
                    for field_name in (
                        "thickness", "milling", "color", "color_source", "milling_source"
                    ):
                        if getattr(part, field_name) is None:
                            setattr(part, field_name, getattr(previous, field_name))
                result.parts.append(part)
                previous_by_block[block_index] = part

    if not result.parts:
        raise PartsImportError(
            "Заголовки нашлись, но ни одной строки с размерами. "
            "Проверьте, что в колонках «Длина» и «Ширина» стоят числа."
        )
    result.parts = _combine_identical_parts(result.parts)
    return result


def _rows_from_xlsx(content: bytes) -> tuple[list[list[object]], str | None]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows: list[list[object]] = []
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index >= MAX_ROWS:
                break
            rows.append(list(row[:MAX_COLUMNS]))
        return rows, sheet.title
    finally:
        workbook.close()


def _rows_from_xls(content: bytes) -> tuple[list[list[object]], str | None]:
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    sheet = book.sheet_by_index(0)
    rows: list[list[object]] = []
    for row_index in range(min(sheet.nrows, MAX_ROWS)):
        rows.append(list(sheet.row_values(row_index))[:MAX_COLUMNS])
    return rows, sheet.name


def _rows_from_csv(content: bytes) -> tuple[list[list[object]], str | None]:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise PartsImportError("Не удалось прочитать файл: неизвестная кодировка.")

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"
    reader = csv.reader(io.StringIO(text), dialect)
    return [row[:MAX_COLUMNS] for _, row in zip(range(MAX_ROWS), reader, strict=False)], None


def parse_parts_file(file_name: str, content: bytes) -> ParseResult:
    """Разбирает .xlsx, .xls или .csv со списком деталей."""
    if not content:
        raise PartsImportError("Файл пустой.")

    extension = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    readers = {
        "xlsx": _rows_from_xlsx,
        "xlsm": _rows_from_xlsx,
        "xls": _rows_from_xls,
        "csv": _rows_from_csv,
    }
    reader = readers.get(extension)
    if reader is None:
        raise PartsImportError(
            f"Формат «{extension or 'без расширения'}» не поддерживается. "
            "Нужен файл .xlsx, .xls или .csv."
        )

    try:
        rows, sheet_name = reader(content)
    except PartsImportError:
        raise
    except Exception as error:  # noqa: BLE001 - показываем причину менеджеру
        raise PartsImportError(f"Файл не читается: {error}") from error

    return parse_rows(rows, sheet_name)
