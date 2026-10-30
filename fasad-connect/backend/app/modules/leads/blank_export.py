"""Бланк заказа в XLSX — тот же лист, что студия печатает на бумаге.

Excel выбран не случайно: заказчик и сейчас ведёт заказы в `.xls`, файл можно
переслать, доправить руками и распечатать откуда угодно. Вёрстка повторяет
бумажный бланк построчно, чтобы в цехе не пришлось переучиваться.
"""

from __future__ import annotations

import io
import math
import re
import textwrap
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from app.modules.leads.models import Lead

SQ_MM_IN_SQ_M = Decimal("1000000")

#: Столько строк в их бумажном бланке. Печатается с запасом — в пустые
#: дописывают позиции от руки.
BLANK_ROWS = 31

#: После адреса сразу начинается таблица. Комментарий печатается под итогами.
HEADER_ROW = 8

HEADERS = ("№", "Длина", "Ширина", "", "М2", "Мм", "Фрезеровка", "Цвет")

COLUMN_WIDTHS = (5, 10, 10, 8, 12, 6, 22, 34)

TITLE = 'Схема на присадку под петли должна быть вид "С ЛИЦА"'
FOOTER_LINES = (
    "Фасад Connect. Бланк заказа.",
    "Проверьте размеры и параметры перед передачей в производство.",
)


def read_parts(lead: Lead) -> list[dict[str, object]]:
    """Позиции заявки — там же, где их читает интерфейс.

    Формат один на всю систему: `parsed_payload_json.raw_payload.parts`. Держать
    второе место чтения нельзя — бумага разошлась бы с экраном.
    """
    payload = lead.parsed_payload_json or {}
    raw = payload.get("raw_payload") or {}
    parts = raw.get("parts")
    if not isinstance(parts, list):
        return []

    rows: list[dict[str, object]] = []
    for item in parts:
        if not isinstance(item, dict):
            continue
        width = item.get("width")
        height = item.get("height")
        if not isinstance(width, (int, float)) or not isinstance(height, (int, float)):
            continue
        quantity = item.get("quantity")
        rows.append(
            {
                "width": int(width),
                "height": int(height),
                "quantity": int(quantity) if isinstance(quantity, (int, float)) and quantity else 1,
                "material": str(item.get("material") or ""),
                "thickness": str(item.get("thickness") or ""),
                "milling": str(item.get("milling") or ""),
            }
        )
    return rows


def part_area_sqm(part: dict[str, object]) -> Decimal:
    return (
        Decimal(str(part["width"])) * Decimal(str(part["height"])) * Decimal(str(part["quantity"]))
    ) / SQ_MM_IN_SQ_M


def _delivery_address(lead: Lead) -> str:
    """Куда везти этот заказ.

    Сначала адрес самой заявки: доставка бывает разовой — на объект, на склад
    монтажников. Если его не назвали, берём адрес карточки клиента: он и есть
    «обычно возим сюда».
    """
    payload = lead.parsed_payload_json or {}
    own = str(payload.get("address") or "").strip()
    if own:
        return own
    return (lead.client.address or "").strip() if lead.client else ""


def _blank_comment(lead: Lead, legacy_comment: str | None) -> str:
    """Только поле «Что нужно клиенту», включая явно оставленное пустым.

    raw_text и старый order.comment содержат автоматически собранную сводку:
    контакт, размеры и расчёт. Это не комментарий менеджера.
    Для старых заявок без отдельного поля сохраняем обычный текст, но не сводку.
    """
    payload = lead.parsed_payload_json or {}
    raw = payload.get("raw_payload") or {}
    if isinstance(raw, dict) and "manual_note" in raw:
        value = raw["manual_note"]
        return value.strip() if isinstance(value, str) else ""
    text = (legacy_comment or "").strip()
    headings = re.findall(r"^(?:Клиент|Телефон|Email|Адрес|Размеры|Всего деталей|Расчёт):", text, re.MULTILINE)
    return "" if len(headings) >= 2 else text


def format_blank_thickness(value) -> str:
    text = "" if value is None else str(value).strip()
    if not re.fullmatch(r"[+-]?\d+(?:[.,]\d+)?", text):
        return text
    text = text.replace(",", ".")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text.replace(".", ",")


def build_blank_xlsx(lead: Lead, comment: str | None, *, order_number: str | None = None,
                     parts_override: list[dict] | None = None, final_amount: Decimal | None = None) -> bytes:
    """Собирает бланк заказа и отдаёт готовый файл байтами.

    Основной источник — raw_payload.manual_note («Что нужно клиенту»).
    ``comment`` используется только для старых заявок без этого поля,
    если он не является автоматически собранной сводкой.
    """
    parts = read_parts(lead) if parts_override is None else parts_override
    # blank-thickness-display-v1: copies for display only
    parts = [{**part, "thickness": format_blank_thickness(part.get("thickness"))} for part in parts]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Бланк заказа"

    thin = Side(style="thin", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    centre = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for index, width in enumerate(COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    last_column = get_column_letter(len(HEADERS))

    sheet.merge_cells(f"A1:{last_column}1")
    sheet["A1"] = TITLE
    sheet["A1"].font = Font(bold=True, italic=True, size=12)
    sheet["A1"].alignment = centre

    sheet.merge_cells(f"A2:{last_column}2")
    sheet["A2"] = f"Бланк заказа: {order_number if order_number is not None else lead.lead_number or ''}"
    sheet["A2"].font = Font(bold=True, size=12)
    sheet["A2"].alignment = centre

    sheet["A4"] = "Дата:"
    sheet["C4"] = lead.created_at.strftime("%d.%m.%Y") if lead.created_at else ""
    sheet["A5"] = "Заказчик"
    sheet["C5"] = lead.client.full_name if lead.client else ""
    sheet["A6"] = "Дата выдачи ориентировочная"
    # Срок печатался пустой подписью без значения, хотя в заявке он есть и его
    # уже назвали клиенту. По бланку без даты в цехе спрашивают у менеджера.
    sheet["C6"] = lead.due_date.strftime("%d.%m.%Y") if lead.due_date else ""
    sheet["E6"] = "Дата выдачи фактическая"
    # Адрес доставки на бумаге: бланк — это лист, который едет с заказом, а
    # адрес до сих пор жил только на экране. Сначала адрес самой заявки (эта
    # поставка может ехать на объект), иначе адрес карточки клиента.
    sheet["A7"] = "Адрес доставки"
    sheet["C7"] = _delivery_address(lead)
    # Адрес длинный: сливаем клетки до края, иначе он обрежется по колонке.
    sheet.merge_cells(f"C7:{last_column}7")
    for cell in ("A4", "A5", "A6", "A7", "E6"):
        sheet[cell].font = Font(bold=False)
    for cell in ("C4", "C5", "C6", "C7"):
        sheet[cell].font = Font(bold=True)

    header_row = HEADER_ROW
    for index, title in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=header_row, column=index, value=title)
        cell.font = Font(bold=True)
        cell.alignment = centre
        cell.border = box

    row_count = max(BLANK_ROWS, len(parts))
    for offset in range(row_count):
        row = header_row + 1 + offset
        part = parts[offset] if offset < len(parts) else None

        values: tuple[object, ...] = (
            offset + 1,
            part["width"] if part else "",
            part["height"] if part else "",
            part["quantity"] if part else "",
            # пустая строка бланка печатает ноль — так в оригинале
            float(part_area_sqm(part)) if part else 0,
            part["thickness"] if part else "",
            part["milling"] if part else "",
            part["material"] if part else "",
        )
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=index, value=value)
            cell.border = box
            cell.alignment = centre
            if index == 5:
                cell.number_format = "0.00"

    totals_row = header_row + 1 + row_count
    total_quantity = sum(int(part["quantity"]) for part in parts)
    total_area = sum((part_area_sqm(part) for part in parts), Decimal("0"))
    if parts:
        sheet.cell(row=totals_row, column=4, value=total_quantity).font = Font(bold=True)
        area_cell = sheet.cell(row=totals_row, column=5, value=float(total_area))
        area_cell.font = Font(bold=True)
        area_cell.number_format = "0.00"

    # сводка по толщинам — в их бланке табличка «16- / 19-» справа от итогов.
    # Правая колонка — площадь в м² по этой толщине, а не количество деталей:
    # цех считает материал квадратами, а не штуками, штуки уже есть в самой
    # таблице выше.
    by_thickness: dict[str, Decimal] = {}
    for part in parts:
        thickness = str(part["thickness"]).strip()
        if thickness:
            by_thickness[thickness] = by_thickness.get(thickness, Decimal("0")) + part_area_sqm(
                part
            )

    thickness_row = totals_row + 1
    for index, key in enumerate(sorted(by_thickness, key=lambda value: (len(value), value))):
        row = thickness_row + index
        left = sheet.cell(row=row, column=7, value=f"{key}-")
        right = sheet.cell(row=row, column=8, value=float(by_thickness[key]))
        right.number_format = "0.00"
        for cell in (left, right):
            cell.border = box
            cell.alignment = centre

    # Комментарий после всей таблицы и сводки толщин, перед подписями.
    # Запас высоты нужен для переноса строк в объединённых ячейках Excel:
    # сам Excel автоматически высоту таких строк не подбирает.
    comment_row = max(totals_row + 2, thickness_row + len(by_thickness))
    comment_text = _blank_comment(lead, comment)
    line_count = sum(max(1, len(textwrap.wrap(line, width=75))) for line in comment_text.splitlines())
    comment_height = max(30, line_count * 15 + 6)
    comment_rows = max(1, math.ceil(comment_height / 390))
    comment_end = comment_row + comment_rows - 1
    sheet.merge_cells(f"A{comment_row}:B{comment_end}")
    sheet.merge_cells(f"C{comment_row}:{last_column}{comment_end}")
    label = sheet.cell(row=comment_row, column=1, value="Комментарий")
    label.alignment = Alignment(horizontal="left", vertical="top")
    value = sheet.cell(row=comment_row, column=3, value=comment_text)
    value.font = Font(bold=True)
    value.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    for row in range(comment_row, comment_end + 1):
        sheet.row_dimensions[row].height = comment_height / comment_rows

    if final_amount is not None:
        amount_row = comment_end + 1
        sheet.cell(row=amount_row, column=1, value="Сумма заказа").font = Font(bold=True)
        sheet.merge_cells(start_row=amount_row, start_column=1, end_row=amount_row, end_column=3)
        amount_cell = sheet.cell(row=amount_row, column=4, value=float(final_amount))
        amount_cell.font = Font(bold=True)
        amount_cell.number_format = '#,##0" ₽"'
    signs_row = comment_end + (3 if final_amount is not None else 2)
    sheet.cell(row=signs_row, column=1, value="Размеры проверил_________________________")
    sheet.cell(row=signs_row + 2, column=1, value="Заказ принял_____________________________")

    footer_row = signs_row + 5
    for index, line in enumerate(FOOTER_LINES):
        row = footer_row + index
        sheet.merge_cells(f"A{row}:{last_column}{row}")
        cell = sheet.cell(row=row, column=1, value=line)
        cell.font = Font(bold=True, italic=True)
        cell.alignment = centre

    # печать под A4 книжной: бланк должен влезать в ширину страницы
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
