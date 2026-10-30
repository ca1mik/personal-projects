"""Бланк заказа в XLSX."""

from __future__ import annotations

import io
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.modules.leads.blank_export import BLANK_ROWS, HEADER_ROW, build_blank_xlsx, read_parts
from app.modules.imports.parser import parse_parts_file


def make_lead(parts: list[dict] | None = None, **overrides):
    payload = {"raw_payload": {"parts": parts if parts is not None else []}}
    defaults = dict(
        id="00000000-0000-0000-0000-000000000001",
        lead_number="З-0042",
        created_at=None,
        due_date=None,
        raw_text="",
        client=SimpleNamespace(full_name="Агат-Сервис", address=None),
        parsed_payload_json=payload,
    )
    defaults.update(overrides)
    return SimpleNamespace(**defaults)


def sheet_of(lead, comment=None):
    return load_workbook(io.BytesIO(build_blank_xlsx(lead, comment))).active


def test_blank_repeats_their_paper_layout():
    lead = make_lead([{"width": 356, "height": 396, "quantity": 5, "thickness": "19"}])
    sheet = sheet_of(lead)

    assert 'вид "С ЛИЦА"' in sheet["A1"].value
    assert sheet["A2"].value == "Бланк заказа: З-0042"
    assert sheet["C5"].value == "Агат-Сервис"
    # строка заголовков таблицы
    assert [sheet.cell(row=HEADER_ROW, column=i).value for i in (1, 2, 3, 5, 6, 7, 8)] == [
        "№",
        "Длина",
        "Ширина",
        "М2",
        "Мм",
        "Фрезеровка",
        "Цвет",
    ]


def test_area_matches_their_formula():
    """М2 = Длина/1000 x Ширина/1000 x Кол-во — 356x396x5 = 0.70488."""
    lead = make_lead([{"width": 356, "height": 396, "quantity": 5}])
    sheet = sheet_of(lead)

    assert sheet.cell(row=HEADER_ROW + 1, column=5).value == pytest.approx(0.70488)
    assert sheet.cell(row=HEADER_ROW + 1, column=5).number_format == "0.00"


def test_all_columns_of_the_blank_are_filled():
    lead = make_lead(
        [
            {
                "width": 391,
                "height": 706,
                "quantity": 2,
                "material": "Эмаль RAL 9003",
                "thickness": "19",
                "milling": "Арт 109",
            }
        ]
    )
    sheet = sheet_of(lead)

    assert [sheet.cell(row=HEADER_ROW + 1, column=i).value for i in (2, 3, 4, 6, 7, 8)] == [
        391,
        706,
        2,
        "19",
        "Арт 109",
        "Эмаль RAL 9003",
    ]


def test_exported_blank_imports_back_with_all_detail_fields():
    lead = make_lead([
        {"width": 746, "height": 396, "quantity": 2, "material": "палома SF 023", "thickness": "16", "milling": "мыло R2"},
        {"width": 746, "height": 596, "quantity": 3, "material": "палома SF 023", "thickness": "16", "milling": "мыло R2"},
    ])

    result = parse_parts_file("blank.xlsx", build_blank_xlsx(lead, None))

    assert [(part.width, part.height, part.quantity) for part in result.parts] == [(746, 396, 2), (746, 596, 3)]
    assert [(part.thickness, part.milling, part.color) for part in result.parts] == [
        ("16", "мыло R2", "палома SF 023"),
        ("16", "мыло R2", "палома SF 023"),
    ]


def test_totals_sum_quantity_and_area():
    lead = make_lead(
        [
            {"width": 1000, "height": 500, "quantity": 2},
            {"width": 500, "height": 500, "quantity": 1},
        ]
    )
    sheet = sheet_of(lead)
    totals = HEADER_ROW + 1 + BLANK_ROWS

    assert sheet.cell(row=totals, column=4).value == 3
    # 1.0 + 0.25 = 1.25 м²
    assert sheet.cell(row=totals, column=5).value == pytest.approx(1.25)


def test_blank_keeps_spare_rows_for_handwriting():
    """Бланк печатают с запасом: пустые строки до 31-й остаются."""
    lead = make_lead([{"width": 300, "height": 300, "quantity": 1}])
    sheet = sheet_of(lead)

    # вторая строка таблицы пустая, но напечатана и содержит ноль в М2
    assert sheet.cell(row=HEADER_ROW + 2, column=2).value in (None, "")
    assert sheet.cell(row=HEADER_ROW + 2, column=5).value == 0
    assert sheet.cell(row=HEADER_ROW + BLANK_ROWS, column=1).value == BLANK_ROWS


# --- отказные сценарии -----------------------------------------------------


def test_lead_without_parts_still_prints_an_empty_blank():
    sheet = sheet_of(make_lead([]))

    assert sheet["A2"].value == "Бланк заказа: З-0042"
    # итогов нет, но пустые строки напечатаны
    assert sheet.cell(row=HEADER_ROW + BLANK_ROWS, column=1).value == BLANK_ROWS


def test_broken_rows_are_skipped_not_crashing():
    parts = [
        {"width": 300, "height": 300, "quantity": 1},
        {"width": "нет", "height": 300, "quantity": 1},
        "мусор",
        {"height": 300},
    ]

    assert len(read_parts(make_lead(parts))) == 1


def test_payload_without_parts_key_is_not_an_error():
    lead = make_lead()
    lead.parsed_payload_json = {"raw_payload": {"quantity": 5}}

    assert read_parts(lead) == []


def test_quantity_defaults_to_one():
    parts = read_parts(make_lead([{"width": 300, "height": 300}]))

    assert parts[0]["quantity"] == 1


@pytest.mark.parametrize('count',[0,1,31,45])
def test_comment_is_below_all_details_and_thickness_totals(count):
    lead=make_lead([{'width':300,'height':400,'quantity':1,'thickness':str(16+i%4)} for i in range(count)])
    comment='Петли слева.\nНе сверлить до согласования.\n'+('Длинный комментарий клиента. '*18)
    sheet=sheet_of(lead,comment)
    row=next(r for r in range(1,sheet.max_row+1) if sheet.cell(r,1).value=='Комментарий')
    assert row>HEADER_ROW+max(BLANK_ROWS,count)
    assert sheet.cell(row,3).value==comment.strip()
    assert sheet.cell(row,3).alignment.wrap_text
    assert sheet.row_dimensions[row].height>=90
    assert all(sheet.cell(r,1).value!='Комментарий' for r in range(1,HEADER_ROW))
    assert all(r<row for r in range(1,sheet.max_row+1) if str(sheet.cell(r,7).value or '').endswith('-'))
    assert next(r for r in range(1,sheet.max_row+1) if str(sheet.cell(r,1).value or '').startswith('Размеры проверил'))>row


@pytest.mark.parametrize('note',['Петли слева. Не сверлить до согласования.','',None])
def test_blank_prints_only_manual_note_not_order_summary(note):
    summary='Клиент: Рифат\nТелефон: +7 960 044 24 93\nАдрес: Богатые Сабы\nРазмеры:\n356x596 — 4 шт.\nВсего деталей: 37\nРасчёт: 72547.26 ₽'
    lead=make_lead([{'width':356,'height':596,'quantity':4}])
    lead.parsed_payload_json['raw_payload']['manual_note']=note
    lead.raw_text=summary
    sheet=sheet_of(lead,summary)
    row=next(r for r in range(1,sheet.max_row+1) if sheet.cell(r,1).value=='Комментарий')
    assert (sheet.cell(row,3).value or '')==(note or '')
    assert sheet.row_dimensions[row].height==30


def test_legacy_automatic_summary_is_not_a_comment():
    sheet=sheet_of(make_lead(), 'Клиент: Рифат\nРазмеры:\n356x596 — 4 шт.')
    row=next(r for r in range(1,sheet.max_row+1) if sheet.cell(r,1).value=='Комментарий')
    assert sheet.cell(row,3).value in ('',None)
