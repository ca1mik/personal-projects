from copy import deepcopy
from decimal import Decimal
import io
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from app.modules.leads.blank_export import build_blank_xlsx, format_blank_thickness


@pytest.mark.parametrize('value,expected', [('19.00','19'),('19','19'),('19,00','19'),
    ('19.50','19,5'), ('19.050','19,05'), ('190.00','190'), ('0.00','0'), (None,''), ('','')])
def test_display_drops_only_fractional_zeroes(value, expected):
    assert format_blank_thickness(value) == expected


def test_xlsx_normalizes_table_and_summary_without_changing_source_or_amount():
    parts = [{'width':100,'height':200,'quantity':1,'thickness':thickness,'milling':'Мыло R2','material':'RAL'}
             for thickness in ['19.00','19','19.50']]
    original = deepcopy(parts)
    lead = SimpleNamespace(parsed_payload_json={}, client=None, created_at=None, due_date=None, lead_number='3205 Доп')
    data = build_blank_xlsx(lead, None, parts_override=parts, order_number='3205-П Доп', final_amount=Decimal('1234'))
    sheet = load_workbook(io.BytesIO(data)).active
    assert [sheet.cell(row=n,column=6).value for n in [9,10,11]] == ['19','19','19,5']
    assert [c.value for row in sheet for c in row].count('19-') == 1
    assert '19,5-' in [c.value for row in sheet for c in row]
    assert 'Бланк заказа: 3205-П Доп' in [c.value for row in sheet for c in row]
    amount_row = next(row for row in sheet if row[0].value == 'Сумма заказа')
    assert amount_row[3].value == 1234
    assert parts == original
