"""Взаиморасчёты с клиентом: баланс, лента поступлений и выписка файлом.

Заказчик ведёт это руками в таблице «сумма / оплата / разница». Здесь
проверяется, что система считает те же три числа и не путает долг с авансом.
"""

from __future__ import annotations

import io

from openpyxl import load_workbook


def _client_with_order(client, auth_headers, *, email, amount="1000.00"):
    """Клиент с одним заказом на заданную сумму."""
    item = client.post(
        "/api/v1/price-book",
        json={
            "article": f"СЧ-{email[:4]}",
            "name": "Позиция для расчётов",
            "kind": "milling",
            "unit": "pcs",
            "price": amount,
        },
        headers=auth_headers,
    ).json()["data"]

    lead = client.post(
        "/api/v1/leads/intake",
        json={
            "source_channel": "manual_manager_input",
            "text": "Заявка для взаиморасчётов",
            "sender": {"full_name": "Расчёты Тест", "phone": "+79990003030", "email": email},
            "raw_payload": {"quantity": 1},
        },
    ).json()["data"]
    order = client.post(
        f"/api/v1/leads/{lead['id']}/create-order", json={}, headers=auth_headers
    ).json()["data"]
    client.put(
        f"/api/v1/orders/{order['id']}/positions",
        json={
            "positions": [
                {
                    "price_item_id": item["id"],
                    "name": "Позиция",
                    "unit": "pcs",
                    "quantity": "1",
                }
            ]
        },
        headers=auth_headers,
    )
    return lead["client_id"], order["id"]


def _statement(client, auth_headers, client_id):
    return client.get(f"/api/v1/clients/{client_id}", headers=auth_headers).json()["data"][
        "statement"
    ]


def test_order_without_payment_is_a_debt(client, auth_headers):
    client_id, _ = _client_with_order(client, auth_headers, email="debt@example.com")

    balance = _statement(client, auth_headers, client_id)["balance"]

    assert balance["charged"] == "1000.00"
    assert balance["paid"] == "0.00"
    assert balance["balance"] == "-1000.00"
    assert balance["status"] == "Долг"
    assert balance["unpaid_orders"] == 1
    assert balance["due_total"] == "1000.00"


def test_full_payment_closes_the_settlement(client, auth_headers):
    client_id, order_id = _client_with_order(client, auth_headers, email="closed@example.com")
    client.post(
        f"/api/v1/orders/{order_id}/payments",
        json={"amount": "1000.00", "method": "cash"},
        headers=auth_headers,
    )

    balance = _statement(client, auth_headers, client_id)["balance"]

    assert balance["balance"] == "0.00"
    assert balance["status"] == "Расчёты закрыты"
    assert balance["unpaid_orders"] == 0


def test_overpayment_becomes_an_advance(client, auth_headers):
    client_id, order_id = _client_with_order(client, auth_headers, email="advance@example.com")
    client.post(
        f"/api/v1/orders/{order_id}/payments",
        json={"amount": "1500.00", "method": "transfer"},
        headers=auth_headers,
    )

    statement = _statement(client, auth_headers, client_id)

    assert statement["balance"]["balance"] == "500.00"
    assert statement["balance"]["status"] == "Аванс"
    # Переплата видна и по самому заказу, а не только в общем балансе.
    assert statement["orders"][0]["overpaid"] == "500.00"
    assert statement["orders"][0]["due"] == "0.00"


def test_ledger_shows_source_and_running_balance(client, auth_headers):
    client_id, order_id = _client_with_order(client, auth_headers, email="ledger@example.com")
    client.post(
        f"/api/v1/orders/{order_id}/payments",
        json={"amount": "400.00", "method": "cash", "comment": "аванс при заказе"},
        headers=auth_headers,
    )
    client.post(
        f"/api/v1/orders/{order_id}/payments",
        json={"amount": "600.00", "method": "account", "comment": "остаток по счёту"},
        headers=auth_headers,
    )

    ledger = _statement(client, auth_headers, client_id)["ledger"]

    # Лента отдаётся свежим сверху: последнее событие — вторая оплата.
    assert [item["kind"] for item in ledger] == ["payment", "payment", "order"]
    assert ledger[0]["source"] == "Расчётный счёт"
    assert ledger[0]["comment"] == "остаток по счёту"
    assert ledger[1]["source"] == "Наличные"
    # Баланс после каждого события: −1000 → −600 → 0
    assert [item["balance_after"] for item in ledger] == ["0.00", "-600.00", "-1000.00"]


def test_order_row_shows_charged_paid_and_due(client, auth_headers):
    client_id, order_id = _client_with_order(client, auth_headers, email="row@example.com")
    client.post(
        f"/api/v1/orders/{order_id}/payments",
        json={"amount": "300.00", "method": "card"},
        headers=auth_headers,
    )

    row = _statement(client, auth_headers, client_id)["orders"][0]

    assert row["charged"] == "1000.00"
    assert row["paid"] == "300.00"
    assert row["due"] == "700.00"
    assert row["payments_count"] == 1


def test_client_without_orders_has_no_settlement(client, auth_headers):
    created = client.post(
        "/api/v1/clients",
        json={"full_name": "Без заказов", "phone": "+79990004040"},
        headers=auth_headers,
    ).json()["data"]

    balance = _statement(client, auth_headers, created["id"])["balance"]

    assert balance["charged"] == "0.00"
    assert balance["status"] == "Нет расчётов"
    assert balance["orders_count"] == 0


def test_statement_downloads_as_xlsx(client, auth_headers):
    client_id, order_id = _client_with_order(client, auth_headers, email="file@example.com")
    client.post(
        f"/api/v1/orders/{order_id}/payments",
        json={"amount": "400.00", "method": "cash", "comment": "аванс"},
        headers=auth_headers,
    )

    response = client.get(f"/api/v1/clients/{client_id}/statement.xlsx", headers=auth_headers)

    assert response.status_code == 200
    assert "filename*=UTF-8''" in response.headers["content-disposition"]

    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Взаиморасчёты", "Лента расчётов"]

    sheet = book["Взаиморасчёты"]
    labels = {row[0].value: row[1].value for row in sheet.iter_rows(min_row=1, max_col=2)}
    assert labels["Начислено"] == 1000
    assert labels["Оплачено"] == 400
    assert labels["Разница"] == -600
    assert labels["Состояние"] == "Долг"

    # Шапку ищем, а не считаем по номеру строки: над таблицей есть заголовок,
    # и жёсткий номер сломается от любой правки вёрстки.
    ledger = book["Лента расчётов"]
    header_row = next(row for row in ledger.iter_rows(values_only=True) if "Откуда" in row)
    assert "Баланс после" in header_row
    assert "Поступило" in header_row


def test_statement_requires_authorization(client, auth_headers):
    client_id, _ = _client_with_order(client, auth_headers, email="auth@example.com")

    assert client.get(f"/api/v1/clients/{client_id}/statement.xlsx").status_code == 401
