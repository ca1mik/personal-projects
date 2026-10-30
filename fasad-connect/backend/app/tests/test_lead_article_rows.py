"""Строка расчёта без размеров не должна пропадать.

Заказчик: «я первый раз заполняю заявку и вручную прописываю артикул, например
128. И всё, этот артикул в памяти должен сохраниться». Пока строка без длины и
ширины выбрасывалась ещё в форме, набранный номер не доезжал до сервера вовсе —
менеджер вводил его, выходил, возвращался, а в подсказках номера нет.

При этом такая строка — не деталь: в раскрой она не идёт и в счётчик фасадов
попадать не должна.
"""

from __future__ import annotations

from decimal import Decimal


def _intake(client, parts: list[dict], email: str):
    return client.post(
        "/api/v1/leads/intake",
        json={
            "source_channel": "manual_manager_input",
            "text": "Расчёт клиенту",
            "sender": {"full_name": "Артикул Тест", "email": email},
            "raw_payload": {"parts": parts},
        },
    ).json()["data"]


def _hints(client, auth_headers) -> dict[str, list[str]]:
    data = client.get("/api/v1/reference-values", headers=auth_headers).json()["data"]
    return {key: [item["value"] for item in data[key]] for key in ("articles", "colors")}


def test_article_without_sizes_is_remembered(client, auth_headers):
    """Главное: набранный номер попадает в память подсказок."""
    _intake(
        client,
        [{"width": None, "height": None, "quantity": 1, "article": "128"}],
        "art@example.com",
    )

    assert "128" in _hints(client, auth_headers)["articles"]


def test_colour_without_sizes_is_remembered_too(client, auth_headers):
    """Цвет копится по тому же правилу: «так же и по цвету, вносим вручную»."""
    _intake(
        client,
        [{"width": None, "height": None, "quantity": 1, "material": "Капучино софт"}],
        "colour@example.com",
    )

    assert "Капучино софт" in _hints(client, auth_headers)["colors"]


def test_row_without_sizes_is_not_counted_as_a_part(client, auth_headers):
    """Строка расчёта — не фасад: в «Деталей» она попадать не должна."""
    lead = _intake(
        client,
        [
            {"width": 716, "height": 396, "quantity": 4, "article": "112"},
            {"width": None, "height": None, "quantity": 1, "article": "128"},
        ],
        "count@example.com",
    )

    summary = client.get(f"/api/v1/leads/{lead['id']}", headers=auth_headers).json()["data"][
        "summary"
    ]

    assert summary["quantity"] == 4, "строка без размеров сосчиталась как деталь"


def test_area_ignores_rows_without_sizes(client, auth_headers):
    """Площадь тоже считается только по фасадам."""
    lead = _intake(
        client,
        [
            {"width": 1000, "height": 500, "quantity": 2, "article": "112"},
            {"width": None, "height": None, "quantity": 1, "article": "128"},
        ],
        "area@example.com",
    )

    summary = client.get(f"/api/v1/leads/{lead['id']}", headers=auth_headers).json()["data"][
        "summary"
    ]

    # 1000 × 500 × 2 = 1 м². Сравниваем числом: точность записи не проверяем.
    assert Decimal(summary["facade_area_sqm"]) == Decimal("1")


def test_such_a_row_does_not_reach_the_printed_blank(client, auth_headers):
    """В бланке печатают детали. Строке без размеров там делать нечего."""
    import io

    from openpyxl import load_workbook

    lead = _intake(
        client,
        [
            {"width": 716, "height": 396, "quantity": 4, "article": "112"},
            {"width": None, "height": None, "quantity": 1, "article": "128"},
        ],
        "blank@example.com",
    )

    response = client.get(f"/api/v1/leads/{lead['id']}/blank.xlsx", headers=auth_headers)
    sheet = load_workbook(io.BytesIO(response.content)).active

    header_row = next(
        row for row in range(1, 20) if str(sheet.cell(row=row, column=1).value or "").strip() == "№"
    )
    # Первая строка таблицы — фасад с размерами, вторая уже пустая заготовка.
    assert sheet.cell(row=header_row + 1, column=2).value == 716
    assert sheet.cell(row=header_row + 2, column=2).value in ("", None)


def test_priced_row_without_sizes_becomes_a_piece_position(client, auth_headers):
    """С ценой такая строка уходит в смету штукой, а не нулём.

    Умножать цену на нулевую площадь значило бы завести позицию за 0 ₽ и молча
    занизить сумму заказа.
    """
    lead = _intake(
        client,
        [
            {
                "width": None,
                "height": None,
                "quantity": 1,
                "article": "128",
                "kind": "other",
                "unit": "pcs",
                "unit_price": "1500.00",
                "amount": "1500.00",
            }
        ],
        "priced@example.com",
    )

    order = client.post(
        f"/api/v1/leads/{lead['id']}/create-order", json={}, headers=auth_headers
    ).json()["data"]

    assert order["final_amount"] == "1500.00"
    position = order["positions"][0]
    assert position["unit"] == "pcs"
    assert position["amount"] == "1500.00"


def test_articles_and_address_are_editable_and_reach_converted_order(client, auth_headers):
    lead = _intake(
        client,
        [
            {"width": 716, "height": 396, "quantity": 2, "unit_price": "1000", "kind": "other"},
            {"width": 800, "height": 400, "quantity": 1, "unit_price": "1000", "kind": "other"},
        ],
        "article-edit@example.com",
    )
    lead_id = lead["id"]
    changed = client.patch(
        f"/api/v1/leads/{lead_id}",
        json={"address": "Москва, ул. Лесная, 5", "part_articles": ["A-101", "B-202"]},
        headers=auth_headers,
    )
    assert changed.status_code == 200, changed.text
    parsed = changed.json()["data"]["parsed_payload_json"]
    assert parsed["address"] == "Москва, ул. Лесная, 5"
    assert [part["article"] for part in parsed["raw_payload"]["parts"]] == ["A-101", "B-202"]

    order_response = client.post(
        f"/api/v1/leads/{lead_id}/create-order", json={}, headers=auth_headers
    )
    assert order_response.status_code == 200, order_response.text
    order = order_response.json()["data"]
    assert [position["article"] for position in order["positions"]] == ["A-101", "B-202"]

    again = client.patch(
        f"/api/v1/leads/{lead_id}",
        json={"part_articles": ["A-303", "B-404"]},
        headers=auth_headers,
    )
    assert again.status_code == 200, again.text
    order_again = client.get(f"/api/v1/orders/{order['id']}", headers=auth_headers)
    assert order_again.status_code == 200, order_again.text
    assert [position["article"] for position in order_again.json()["data"]["positions"]] == [
        "A-303", "B-404"
    ]


def test_article_edit_rejects_missing_rows(client, auth_headers):
    lead = _intake(
        client,
        [{"width": 716, "height": 396, "quantity": 2}],
        "article-invalid@example.com",
    )
    response = client.patch(
        f"/api/v1/leads/{lead['id']}",
        json={"part_articles": []},
        headers=auth_headers,
    )
    assert response.status_code == 422
