"""Бланк заказа файлом.

Это лист, который едет с заказом в цех и к клиенту. Всё, что нужно знать при
отгрузке, должно быть на нём: раньше не было ни адреса доставки, ни срока —
подпись «Дата выдачи ориентировочная» печаталась пустой, хотя дата в заявке
есть и её уже назвали клиенту.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook


def _intake(client, *, sender: dict | None = None, **extra):
    payload = {
        "source_channel": "manual_manager_input",
        "text": "Фасады на кухню",
        "sender": sender or {"full_name": "Бланк Тест", "phone": "+79170009000"},
        "raw_payload": {
            "parts": [{"width": 716, "height": 396, "quantity": 4, "material": "RAL 9003"}]
        },
    }
    payload.update(extra)
    return client.post("/api/v1/leads/intake", json=payload).json()["data"]


def _sheet(client, auth_headers, lead_id: str):
    response = client.get(f"/api/v1/leads/{lead_id}/blank.xlsx", headers=auth_headers)
    assert response.status_code == 200, response.text
    return load_workbook(io.BytesIO(response.content)).active


def _cells(sheet) -> dict[str, object]:
    """Значения шапки по подписи слева: так тест не привязан к номеру строки."""
    found: dict[str, object] = {}
    for row in range(1, sheet.max_row + 1):
        label = sheet.cell(row=row, column=1).value
        if isinstance(label, str) and label.strip():
            found[label.strip()] = sheet.cell(row=row, column=3).value
    return found


def test_blank_carries_the_delivery_address(client, auth_headers):
    """Адрес доставки печатается — по нему и везут."""
    lead = _intake(
        client,
        sender={
            "full_name": "Ольга",
            "phone": "+79170009001",
            "address": "Казань, Восстания 100, склад",
        },
    )

    header = _cells(_sheet(client, auth_headers, lead["id"]))

    assert header["Адрес доставки"] == "Казань, Восстания 100, склад"


def test_blank_falls_back_to_the_client_card_address(client, auth_headers):
    """Адрес этой поставки не назвали — печатаем «куда обычно возим»."""
    existing = client.post(
        "/api/v1/clients",
        json={"full_name": "Агат", "phone": "+79170009002", "address": "Казань, Рахимова 8"},
        headers=auth_headers,
    ).json()["data"]
    lead = _intake(client, sender={"full_name": "Агат", "client_id": existing["id"]})

    header = _cells(_sheet(client, auth_headers, lead["id"]))

    assert header["Адрес доставки"] == "Казань, Рахимова 8"


def test_blank_without_any_address_prints_empty(client, auth_headers):
    """Адреса нет нигде — пустая клетка, а не выдуманный адрес."""
    lead = _intake(client)

    header = _cells(_sheet(client, auth_headers, lead["id"]))

    assert header["Адрес доставки"] in ("", None)


def test_blank_carries_the_comment(client, auth_headers):
    """Комментарий из заявки печатается на бланке — раньше для него не было
    ни ячейки, ни чтения этого поля, и цех получал лист без того, что клиент
    просил дополнительно."""
    lead = _intake(client, text="Кромка везде ПВХ, петли уже есть свои — не докупать")

    sheet = _sheet(client, auth_headers, lead["id"])
    assert _cells(sheet)["Комментарий"] == "Кромка везде ПВХ, петли уже есть свои — не докупать"
    comment_row=next(row for row in range(1,sheet.max_row+1) if sheet.cell(row,1).value=="Комментарий")
    last_part_row=max(row for row in range(1,sheet.max_row+1) if isinstance(sheet.cell(row,1).value,int))
    assert comment_row>last_part_row


def test_blank_carries_the_due_date(client, auth_headers):
    """Срок печатается рядом со своей подписью, а не остаётся пустым."""
    lead = _intake(client, due_date="2026-09-15")

    header = _cells(_sheet(client, auth_headers, lead["id"]))

    assert header["Дата выдачи ориентировочная"] == "15.09.2026"


@pytest.mark.parametrize('converted',[False,True])
@pytest.mark.parametrize('note',['Петли слева. Только это пожелание.',None])
def test_export_uses_only_what_customer_needs_field(client,auth_headers,converted,note):
    lead=_intake(client,text='Клиент: Рифат\nТелефон: +79600442493\nРазмеры:\n356x596 — 4 шт.\nРасчёт: 72547.26 ₽',raw_payload={
        'parts':[{'width':356,'height':596,'quantity':4}], 'manual_note':note,
    })
    if converted:
        response=client.post(f"/api/v1/leads/{lead['id']}/create-order",json={},headers=auth_headers)
        assert response.status_code==200,response.text
    sheet=_sheet(client,auth_headers,lead['id'])
    assert (_cells(sheet)['Комментарий'] or '')==(note or '')


def test_blank_thickness_summary_shows_area_not_count(client, auth_headers):
    """Табличка «16- / 19-» справа от итогов — площадь по толщине, не штуки.

    Раньше правая колонка печатала количество деталей этой толщины — то же
    самое, что уже видно в основной таблице. Цех заказывает материал
    квадратами, а не штуками, поэтому там должна быть площадь.
    """
    lead = _intake(
        client,
        raw_payload={
            "parts": [
                {"width": 1000, "height": 500, "quantity": 2, "thickness": 16, "material": "Белый"},
                {"width": 600, "height": 400, "quantity": 3, "thickness": 16, "material": "Белый"},
            ]
        },
    )
    sheet = _sheet(client, auth_headers, lead["id"])

    row = next(
        row for row in range(1, 60) if str(sheet.cell(row=row, column=7).value or "") == "16-"
    )
    area = sheet.cell(row=row, column=8).value

    # 1000*500*2/1e6 + 600*400*3/1e6 = 1.0 + 0.72 = 1.72 м² — не 5 (штуки).
    assert area == pytest.approx(1.72, abs=0.001)


def test_blank_still_lists_the_parts(client, auth_headers):
    """Шапка выросла на строку — таблица деталей от этого не должна съехать."""
    lead = _intake(client)
    sheet = _sheet(client, auth_headers, lead["id"])

    header_row = next(
        row for row in range(1, 20) if str(sheet.cell(row=row, column=1).value or "").strip() == "№"
    )
    first = sheet.cell(row=header_row + 1, column=1).value
    width = sheet.cell(row=header_row + 1, column=2).value

    assert first == 1
    assert width == 716
