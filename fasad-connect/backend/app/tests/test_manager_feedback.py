"""Customer regressions: whole order sums, full XLSX numbers, PVC and moved files."""
import io
import uuid
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.core.exceptions import NotFoundError
from app.modules.directions.models import ProductionDirection
from app.modules.documents.models import Document
from app.modules.documents.service import DocumentService
from app.modules.orders.models import Order
from app.modules.orders.money import round_order_amount
from app.storage.local import LocalStorageBackend
from app.tests.test_order_lifecycle_trash import make_order


@pytest.mark.parametrize("raw,expected", [("13764.09", "13764.00"), ("10.50", "10.00"), ("10.51", "11.00"), ("10.49", "10.00"), ("10.99", "11.00"), ("0", "0.00")])
def test_whole_roubles(raw, expected):
    assert str(round_order_amount(raw)) == expected


def test_old_kopecks_no_longer_make_debt_and_payments_stay_exact(client, auth_headers, db_session):
    _, order = make_order(client, auth_headers, "4690")
    stored = db_session.get(Order, uuid.UUID(order["id"]))
    stored.final_amount = Decimal("13764.09")
    db_session.commit()
    created = stored.created_at
    path = f"/api/v1/orders/{order['id']}"
    paid = client.post(path + "/payments", json={"amount":"13764.00", "method":"cash"}, headers=auth_headers)
    assert paid.status_code == 200, paid.text
    current = client.get(path, headers=auth_headers).json()["data"]
    assert current["final_amount"] == "13764.00"
    assert current["balance"]["balance"] == "0.00"
    assert current["balance"]["status"] == "Оплачен полностью"
    assert current["payments"][0]["amount"] == "13764.00"
    db_session.refresh(stored)
    assert stored.created_at == created


def test_blank_has_current_direction_letter_and_additional_number(client, auth_headers, db_session):
    db_session.add(ProductionDirection(code="pvc", name="ПВХ", number_suffix="П"))
    db_session.commit()
    lead, order = make_order(client, auth_headers, "4690")
    path = f"/api/v1/orders/{order['id']}"
    renamed = client.patch(path, json={"order_number":"004690", "direction_code":"pvc"}, headers=auth_headers)
    assert renamed.status_code == 200, renamed.text
    def number(order_id=None):
        response = client.get(f"/api/v1/leads/{lead['id']}/blank.xlsx" + (f"?orderId={order_id}" if order_id else ""), headers=auth_headers)
        assert response.status_code == 200, response.text
        from urllib.parse import unquote
        title = load_workbook(io.BytesIO(response.content)).active["A2"].value
        filename = unquote(response.headers["content-disposition"].split("filename*=UTF-8''")[1])
        assert filename == title.removeprefix("Бланк заказа: ").replace("/", "-") + ".xlsx"
        return title
    assert number() == "Бланк заказа: 004690-П"
    assert number(order["id"]) == "Бланк заказа: 004690-П"
    extra = client.post(path + "/additional", json={}, headers=auth_headers)
    assert extra.status_code == 200, extra.text
    assert number(extra.json()["data"]["id"]) == "Бланк заказа: " + extra.json()["data"]["order_number"]


def test_enamel_cannot_save_film_but_pvc_can(client, auth_headers, db_session):
    for code,suffix in [("pvc","П"),("enamel", "")]:
        db_session.add(ProductionDirection(code=code, name=code, number_suffix=suffix))
    db_session.commit()
    _, order = make_order(client, auth_headers, "10")
    path = f"/api/v1/orders/{order['id']}"
    client.patch(path, json={"direction_code":"enamel"}, headers=auth_headers)
    response = client.post(path + "/save-cutting-layout", json={"sheet_settings":{"workKind":"film"}}, headers=auth_headers)
    assert response.status_code == 422 and "ПВХ" in response.json()["error"]["message"]
    client.patch(path, json={"direction_code":"pvc"}, headers=auth_headers)
    # Empty snapshots are accepted just as before; direction is the eligibility gate.
    response = client.post(path + "/save-cutting-layout", json={"sheet_settings":{"workKind":"film"}}, headers=auth_headers)
    assert response.status_code == 200, response.text


def test_file_can_be_read_after_portable_folder_moves(tmp_path, monkeypatch):
    order_id = uuid.uuid4()
    storage = LocalStorageBackend(tmp_path)
    key = f"documents/{order_id}/attachment/test.pdf"
    storage.save_bytes(key, b"%PDF old attachment")
    monkeypatch.setattr("app.modules.documents.service.get_storage_backend", lambda: storage)
    doc = Document(order_id=order_id, file_path="C:\\old-fasad\\storage\\" + key.replace("/", "\\"), file_name="test.pdf", doc_type="attachment")
    assert DocumentService.read_content(doc) == b"%PDF old attachment"
    doc.file_path = f"storage/documents/{order_id}/../missing.pdf"
    with pytest.raises(NotFoundError):
        DocumentService.read_content(doc)
    doc.file_path = f"documents/{order_id}/missing.pdf"
    with pytest.raises(NotFoundError, match="не найден"):
        DocumentService.read_content(doc)


@pytest.mark.parametrize("price,agreed", [("100.00", "700.00"), ("0.00", "333.00"), ("100.00", "0.00")])
def test_agreed_amount_is_independent_of_price_and_survives_estimate_save(client, auth_headers, price, agreed):
    lead = client.post("/api/v1/leads/intake", json={
        "source_channel":"manual_manager_input", "text":"Независимая стоимость",
        "sender":{"full_name":"Цена отдельно", "phone":"+79990007000"},
        "raw_payload":{"parts":[{"width":1000,"height":1000,"quantity":1,"unit":"sqm","unit_price":price,"amount":agreed,"amount_override":agreed}]},
    }).json()["data"]
    response = client.post(f"/api/v1/leads/{lead['id']}/create-order", json={}, headers=auth_headers)
    assert response.status_code == 200, response.text
    order = response.json()["data"]
    assert len(order["positions"]) == 1
    assert order["positions"][0]["unit_price"] == price
    assert order["positions"][0]["amount"] == agreed
    assert order["final_amount"] == agreed
    result = client.put(f"/api/v1/orders/{order['id']}/positions", json={"positions":[{
        "is_custom":True,"name":"Согласованная сумма","unit":"sqm","width":"1000","height":"1000","quantity":"1",
        "unit_price":price,"amount_override":agreed,
    }]}, headers=auth_headers)
    assert result.status_code == 200, result.text
    assert result.json()["data"][0]["unit_price"] == price
    assert result.json()["data"][0]["amount"] == agreed
