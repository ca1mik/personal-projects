"""Editing an order updates its summary, blank, totals and printable layout input."""
import io
import json
import uuid
from datetime import timedelta

import pytest
from openpyxl import load_workbook

from app.modules.cutting_export.models import Export
from app.storage.factory import get_storage_backend


def setup_order(client, headers):
    lead = client.post('/api/v1/leads/intake', json={
        'source_channel':'manual_manager_input',
        'text':'Клиент: Синхронизация\nРазмеры:\n123x421 — 2 шт.\nРасчёт: 999 ₽\nНе потерять примечание',
        'sender':{'full_name':'Синхронизация', 'phone':'+79998881234'},
        'raw_payload':{'manual_note':'Петли слева', 'parts':[{
            'width':123,'height':421,'quantity':2,'thickness':2,'material':'Старый цвет'}]},
    }).json()['data']
    response = client.post(f"/api/v1/leads/{lead['id']}/create-order",json={},headers=headers)
    assert response.status_code == 200, response.text
    return lead, response.json()['data']


def save(client, headers, order_id, *, width=1000, quantity=3, price=100, empty=False):
    row = {'is_custom':True,'name':'Фасад','kind':'other','unit':'sqm','unit_price':str(price),
           'width':str(width),'height':'500','quantity':str(quantity),'thickness':'19',
           'milling':'Фреза новая','color':'Новый цвет'}
    response=client.put(f'/api/v1/orders/{order_id}/positions',headers=headers,
                        json={'positions':[] if empty else [row]})
    assert response.status_code==200,response.text
    return response.json()['data']


def sheet(client, headers, lead_id, order_id):
    response=client.get(f'/api/v1/leads/{lead_id}/blank.xlsx?orderId={order_id}',headers=headers)
    assert response.status_code==200,response.text
    return load_workbook(io.BytesIO(response.content)).active


def test_changed_dimensions_reach_summary_xlsx_and_money_without_rewriting_lead(client,auth_headers):
    lead,order=setup_order(client,auth_headers)
    save(client,auth_headers,order['id'])
    current=client.get(f"/api/v1/orders/{order['id']}",headers=auth_headers).json()['data']
    assert current['final_amount']=='150.00'
    assert current['specification_parts'][0]['width']==1000
    assert '1000x500 — 3 шт., 19 мм, Новый цвет' in current['comment']
    assert '123x421' not in current['comment']
    assert 'Расчёт: 150.00 ₽' in current['comment']
    assert 'Петли слева' in current['comment']
    book=sheet(client,auth_headers,lead['id'],order['id'])
    assert [book.cell(9,col).value for col in range(2,9)]==[1000,500,3,1.5,'19.00','Фреза новая','Новый цвет']
    amount_row=next(row for row in book.iter_rows() if row[0].value=='Сумма заказа')
    assert amount_row[3].value==150
    original=client.get(f"/api/v1/leads/{lead['id']}",headers=auth_headers).json()['data']
    assert original['parsed_payload_json']['raw_payload']['parts'][0]['width']==123
    client_card=client.get(f"/api/v1/clients/{order['client_id']}",headers=auth_headers).json()['data']
    assert client_card['order_history'][0]['final_amount']=='150.00'


def test_cleared_estimate_does_not_restore_original_lead_dimensions(client,auth_headers):
    lead,order=setup_order(client,auth_headers)
    save(client,auth_headers,order['id'],empty=True)
    current=client.get(f"/api/v1/orders/{order['id']}",headers=auth_headers).json()['data']
    assert current['specification_parts']==[]
    assert current['final_amount']=='0.00'
    assert 'Всего деталей: 0' in current['comment']
    book=sheet(client,auth_headers,lead['id'],order['id'])
    assert book['B9'].value is None


def test_xlsx_keeps_fractional_dimensions_and_their_exact_area(client,auth_headers):
    lead,order=setup_order(client,auth_headers)
    save(client,auth_headers,order['id'],width=1000.5)
    book=sheet(client,auth_headers,lead['id'],order['id'])
    assert book['B9'].value==1000.5
    assert book['E9'].value==pytest.approx(1.50075)


def test_additional_order_blank_is_independent_of_parent(client,auth_headers):
    lead,order=setup_order(client,auth_headers)
    save(client,auth_headers,order['id'],width=900)
    response=client.post(f"/api/v1/orders/{order['id']}/additional",json={},headers=auth_headers)
    assert response.status_code==200,response.text
    extra=response.json()['data']
    assert sheet(client,auth_headers,lead['id'],extra['id'])['B9'].value is None
    save(client,auth_headers,extra['id'],width=450)
    assert sheet(client,auth_headers,lead['id'],extra['id'])['B9'].value==450
    assert sheet(client,auth_headers,lead['id'],order['id'])['B9'].value==900


@pytest.mark.parametrize('change', ['dimensions','quantity','price'])
def test_saved_layout_invalidated_by_geometry_only(client,auth_headers,change):
    _,order=setup_order(client,auth_headers)
    save(client,auth_headers,order['id'])
    path=f"/api/v1/orders/{order['id']}"
    response=client.post(path+'/save-cutting-layout',headers=auth_headers,
                         json={'sheet_settings':{'workKind':'cutting'},'sheets':[]})
    assert response.status_code==200,response.text
    assert client.get(path+'/cutting-layout/latest?work_kind=cutting',headers=auth_headers).json()['data'] is not None
    save(client,auth_headers,order['id'],width=800 if change=='dimensions' else 1000,
         quantity=4 if change=='quantity' else 3,price=200 if change=='price' else 100)
    snapshot=client.get(path+'/cutting-layout/latest?work_kind=cutting',headers=auth_headers).json()['data']
    assert (snapshot is not None)==(change=='price')


def test_legacy_snapshot_saved_before_estimate_change_is_not_reused(client,auth_headers,db_session):
    _,order=setup_order(client,auth_headers)
    path=f"/api/v1/orders/{order['id']}"
    response=client.post(path+'/save-cutting-layout',headers=auth_headers,json={'sheet_settings':{'workKind':'cutting'},'sheets':[]})
    export=db_session.get(Export,uuid.UUID(response.json()['data']['id']))
    storage=get_storage_backend()
    snapshot=json.loads(storage.read_bytes(export.file_path))
    snapshot.pop('specification_signature')
    storage.save_text(export.file_path,json.dumps(snapshot))
    export.created_at-=timedelta(minutes=1)
    db_session.commit()
    save(client,auth_headers,order['id'])
    assert client.get(path+'/cutting-layout/latest?work_kind=cutting',headers=auth_headers).json()['data'] is None
