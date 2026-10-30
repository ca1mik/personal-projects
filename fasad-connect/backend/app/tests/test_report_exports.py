"""Отчёты файлом: цифры в Excel должны совпадать с тем, что на экране.

Две ветки расчёта разошлись бы на первом же округлении, поэтому выгрузка берёт
данные из тех же сервисов, что рисуют экран, — а тест это сторожит.
"""

from __future__ import annotations

import io

from openpyxl import load_workbook

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _priced_order(client, auth_headers, *, email, article, price, width, height, quantity):
    item = client.post(
        "/api/v1/price-book",
        json={
            "article": article,
            "name": f"Позиция {article}",
            "kind": "milling",
            "unit": "sqm",
            "price": price,
        },
        headers=auth_headers,
    ).json()["data"]
    lead = client.post(
        "/api/v1/leads/intake",
        json={
            "source_channel": "manual_manager_input",
            "text": "Заявка для отчёта",
            "sender": {"full_name": "Отчёт Тест", "phone": "+79990005050", "email": email},
            "raw_payload": {"quantity": quantity},
        },
    ).json()["data"]
    order = client.post(
        f"/api/v1/leads/{lead['id']}/create-order", json={}, headers=auth_headers
    ).json()["data"]
    client.put(
        f"/api/v1/orders/{order['id']}/positions",
        json={
            "positions": [
                {
                    "price_item_id": item["id"],
                    "name": "Фасады",
                    "width": str(width),
                    "height": str(height),
                    "quantity": str(quantity),
                }
            ]
        },
        headers=auth_headers,
    )
    return order["id"]


def _sheet_values(content: bytes, sheet_name: str | None = None):
    book = load_workbook(io.BytesIO(content))
    sheet = book[sheet_name] if sheet_name else book.active
    return book, sheet, [row for row in sheet.iter_rows(values_only=True)]


def test_volume_report_xlsx_matches_the_screen(client, auth_headers):
    _priced_order(
        client,
        auth_headers,
        email="vol@example.com",
        article="ОТЧ-1",
        price="1000.00",
        width=1000,
        height=1000,
        quantity=1,
    )

    on_screen = client.get("/api/v1/reports/volume?period=all", headers=auth_headers).json()["data"]
    response = client.get("/api/v1/reports/volume.xlsx?period=all", headers=auth_headers)

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(XLSX)
    assert "filename*=UTF-8''" in response.headers["content-disposition"]

    book, _, rows = _sheet_values(response.content)
    assert book.sheetnames == ["Объёмы", "По месяцам"]

    labels = {row[0]: row[1] for row in rows if row and row[0]}
    assert labels["Заказов"] == on_screen["total"]["orders_count"]
    assert labels["Сумма"] == float(on_screen["total"]["amount"])


def test_work_kind_report_xlsx_has_both_sheets(client, auth_headers):
    _priced_order(
        client,
        auth_headers,
        email="kind@example.com",
        article="ОТЧ-2",
        price="1200.00",
        width=1000,
        height=1000,
        quantity=2,
    )

    on_screen = client.get("/api/v1/reports/work-kinds?period=all", headers=auth_headers).json()[
        "data"
    ]
    response = client.get("/api/v1/reports/work-kinds.xlsx?period=all", headers=auth_headers)

    book, _, rows = _sheet_values(response.content)
    assert book.sheetnames == ["Обработка", "По артикулам"]

    labels = {row[0]: row[1] for row in rows if row and row[0]}
    assert labels["Сумма"] == float(on_screen["total"]["amount"])

    # Итоговая строка таблицы подписана и содержит ту же сумму.
    total_row = next(row for row in rows if row and row[0] == "Итого")
    assert total_row[5] == float(on_screen["total"]["amount"])


def test_money_is_a_number_not_text(client, auth_headers):
    """Столбец сумм должен складываться в Excel формулой."""
    _priced_order(
        client,
        auth_headers,
        email="num@example.com",
        article="ОТЧ-3",
        price="500.00",
        width=1000,
        height=1000,
        quantity=1,
    )

    response = client.get("/api/v1/reports/work-kinds.xlsx?period=all", headers=auth_headers)
    book = load_workbook(io.BytesIO(response.content))
    sheet = book["Обработка"]

    # Ищем шапку таблицы, а не строку сводки: «Сумма» встречается и там, и там.
    header_row = next(
        row for row in sheet.iter_rows() if any(cell.value == "Тип работы" for cell in row)
    )
    column = next(cell.column for cell in header_row if cell.value == "Сумма")
    value_cell = sheet.cell(row=header_row[0].row + 1, column=column)

    assert isinstance(value_cell.value, (int, float))
    assert value_cell.number_format == "# ##0.00"


def test_clients_report_xlsx(client, auth_headers):
    _priced_order(
        client,
        auth_headers,
        email="cl@example.com",
        article="ОТЧ-4",
        price="900.00",
        width=1000,
        height=1000,
        quantity=1,
    )

    response = client.get("/api/v1/reports/clients.xlsx?period=all", headers=auth_headers)

    assert response.status_code == 200
    book, sheet, rows = _sheet_values(response.content)
    header_row = next(row for row in rows if row and "Клиент" in row)
    assert "Долг" not in header_row  # это отчёт по клиентам, не выписка
    assert "Сумма" in header_row


def test_volume_report_area_reflects_smeta_positions(client, auth_headers):
    """Регрессия: площадь считалась только по каталожным ``OrderItem``.

    Заказы из актуального флоу (лид → заказ → смета) собраны ``OrderPosition``,
    а не каталогом, и отчёт молча показывал 0,00000 м² при настоящей сумме и
    количестве — заказчик увидел это на «Клиенты за период».
    """
    _priced_order(
        client,
        auth_headers,
        email="area-vol@example.com",
        article="ОТЧ-AREA-1",
        price="1000.00",
        width=1000,
        height=2000,
        quantity=3,
    )

    on_screen = client.get("/api/v1/reports/volume?period=all", headers=auth_headers).json()["data"]
    # 1000 x 2000 x 3 / 1_000_000 = 6.0 м²
    assert float(on_screen["total"]["facade_area_sqm"]) == 6.0


def test_clients_report_area_reflects_smeta_positions(client, auth_headers):
    _priced_order(
        client,
        auth_headers,
        email="area-cl@example.com",
        article="ОТЧ-AREA-2",
        price="1000.00",
        width=500,
        height=2000,
        quantity=2,
    )

    on_screen = client.get("/api/v1/reports/clients?period=all", headers=auth_headers).json()[
        "data"
    ]
    row = next(item for item in on_screen["rows"] if item["email"] == "area-cl@example.com")
    # 500 x 2000 x 2 / 1_000_000 = 2.0 м²
    assert float(row["facade_area_sqm"]) == 2.0
    assert float(on_screen["total"]["facade_area_sqm"]) == 2.0


def test_orders_report_area_reflects_smeta_positions(client, auth_headers):
    _priced_order(
        client,
        auth_headers,
        email="area-ord@example.com",
        article="ОТЧ-AREA-3",
        price="1000.00",
        width=2000,
        height=1000,
        quantity=1,
    )

    on_screen = client.get("/api/v1/reports/orders?period=all", headers=auth_headers).json()["data"]
    row = next(item for item in on_screen["rows"] if item["client_name"] == "Отчёт Тест")
    # 2000 x 1000 x 1 / 1_000_000 = 2.0 м²
    assert float(row["facade_area_sqm"]) == 2.0


def test_production_time_report_xlsx(client, auth_headers):
    response = client.get("/api/v1/reports/production-time.xlsx?period=all", headers=auth_headers)

    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Этапы", "Заказы"]


def test_empty_period_still_returns_a_file(client, auth_headers):
    """Пустой период — не ошибка: файл открывается, просто без строк."""
    response = client.get(
        "/api/v1/reports/volume.xlsx?period=custom&date_from=2000-01-01&date_to=2000-01-02",
        headers=auth_headers,
    )

    assert response.status_code == 200
    book, _, rows = _sheet_values(response.content)
    labels = {row[0]: row[1] for row in rows if row and row[0]}
    assert labels["Заказов"] == 0


def test_report_exports_require_authorization(client):
    for path in (
        "/api/v1/reports/volume.xlsx",
        "/api/v1/reports/work-kinds.xlsx",
        "/api/v1/reports/clients.xlsx",
        "/api/v1/reports/orders.xlsx",
        "/api/v1/reports/production-time.xlsx",
    ):
        assert client.get(path).status_code == 401


# --- отбор направления доходит до файла -------------------------------------


def _direction(db_session, code: str, name: str, suffix: str, factor: str = "1.000"):
    from decimal import Decimal

    from app.modules.directions.models import ProductionDirection

    row = ProductionDirection(
        code=code,
        name=name,
        number_suffix=suffix,
        material_area_factor=Decimal(factor),
        sort_order=1,
        is_active=True,
    )
    db_session.add(row)
    db_session.commit()
    return row


def _order_of_direction(client, auth_headers, code: str, email: str, price: str):
    """Заказ нужного направления со сметой — так он попадает во все три отчёта."""
    lead = client.post(
        "/api/v1/leads/intake",
        json={
            "source_channel": "manual_manager_input",
            "direction_code": code,
            "text": "Заявка для выгрузки",
            "sender": {"full_name": f"Клиент {code or 'без'}", "email": email},
            "raw_payload": {},
        },
    ).json()["data"]
    order = client.post(
        f"/api/v1/leads/{lead['id']}/create-order", json={}, headers=auth_headers
    ).json()["data"]
    item = client.post(
        "/api/v1/price-book",
        json={
            "article": f"ВЫГ-{code or 'нет'}",
            "name": f"Работа {code or 'без направления'}",
            "kind": "milling",
            "unit": "pcs",
            "price": price,
        },
        headers=auth_headers,
    ).json()["data"]
    client.put(
        f"/api/v1/orders/{order['id']}/positions",
        json={
            "positions": [
                {"price_item_id": item["id"], "name": "Позиция", "quantity": "1", "unit": "pcs"}
            ]
        },
        headers=auth_headers,
    )
    return order


def _pairs(rows):
    """Шапка файла: «подпись → значение» из первых колонок.

    Берём первое вхождение: ниже идёт таблица, и её заголовок «Направление |
    Заказов» перетирал бы подпись из шапки.
    """
    found: dict = {}
    for row in rows:
        if row and row[0] and row[0] not in found:
            found[row[0]] = row[1]
    return found


def test_volume_xlsx_respects_direction_and_prints_it(client, auth_headers, db_session):
    """Цифры в файле совпадают с экраном, и в шапке написано направление.

    Регрессия: `volume.xlsx` не принимал `direction_code` вовсе, поэтому файл
    считался по всему производству, пока на экране был отобран шпон.
    """
    _direction(db_session, "veneer", "Шпон", "Ш")
    _direction(db_session, "pvc", "ПВХ", "П")
    _order_of_direction(client, auth_headers, "veneer", "vn-x@example.com", "1000.00")
    _order_of_direction(client, auth_headers, "pvc", "pvc-x@example.com", "2000.00")

    on_screen = client.get(
        "/api/v1/reports/volume?period=all&direction_code=veneer", headers=auth_headers
    ).json()["data"]
    response = client.get(
        "/api/v1/reports/volume.xlsx?period=all&direction_code=veneer", headers=auth_headers
    )

    assert response.status_code == 200
    _, _, rows = _sheet_values(response.content)
    labels = _pairs(rows)
    assert labels["Направление"] == "Шпон"
    assert labels["Сумма"] == float(on_screen["total"]["amount"]) == 1000.0


def test_work_kinds_xlsx_respects_direction(client, auth_headers, db_session):
    """`work-kinds.xlsx` тоже считает только по выбранному направлению."""
    _direction(db_session, "veneer", "Шпон", "Ш")
    _direction(db_session, "pvc", "ПВХ", "П")
    _order_of_direction(client, auth_headers, "veneer", "vn-w@example.com", "1000.00")
    _order_of_direction(client, auth_headers, "pvc", "pvc-w@example.com", "2000.00")

    response = client.get(
        "/api/v1/reports/work-kinds.xlsx?period=all&direction_code=pvc", headers=auth_headers
    )

    _, _, rows = _sheet_values(response.content)
    labels = _pairs(rows)
    assert labels["Направление"] == "ПВХ"
    assert labels["Сумма"] == 2000.0


def test_production_time_xlsx_respects_direction(client, auth_headers, db_session):
    """`production-time.xlsx` принимает отбор и печатает его в шапке."""
    _direction(db_session, "veneer", "Шпон", "Ш")

    response = client.get(
        "/api/v1/reports/production-time.xlsx?period=all&direction_code=veneer",
        headers=auth_headers,
    )

    assert response.status_code == 200
    _, _, rows = _sheet_values(response.content)
    assert _pairs(rows)["Направление"] == "Шпон"


def test_xlsx_without_direction_says_all(client, auth_headers, db_session):
    """Без отбора в шапке честное «Все направления», а не пустая строка."""
    _direction(db_session, "veneer", "Шпон", "Ш")
    _order_of_direction(client, auth_headers, "veneer", "vn-all@example.com", "1000.00")

    response = client.get("/api/v1/reports/volume.xlsx?period=all", headers=auth_headers)

    _, _, rows = _sheet_values(response.content)
    assert _pairs(rows)["Направление"] == "Все направления"


def test_unknown_direction_gives_an_empty_file(client, auth_headers, db_session):
    """Неизвестный код — пустой файл с этим кодом в шапке, а не выгрузка по всем."""
    _direction(db_session, "veneer", "Шпон", "Ш")
    _order_of_direction(client, auth_headers, "veneer", "vn-none@example.com", "1000.00")

    response = client.get(
        "/api/v1/reports/volume.xlsx?period=all&direction_code=такого-нет", headers=auth_headers
    )

    assert response.status_code == 200
    _, _, rows = _sheet_values(response.content)
    labels = _pairs(rows)
    assert labels["Направление"] == "такого-нет"
    assert labels["Заказов"] == 0


# --- клиенты: два листа, порог и галочки ------------------------------------


def test_clients_xlsx_has_two_sheets(client, auth_headers):
    """Обещанные в docs два листа: сводка и заказы этих же клиентов."""
    _priced_order(
        client,
        auth_headers,
        email="two@example.com",
        article="ОТЧ-Л",
        price="900.00",
        width=1000,
        height=1000,
        quantity=1,
    )

    response = client.get("/api/v1/reports/clients.xlsx?period=all", headers=auth_headers)

    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Клиенты", "Заказы"]

    detail = [row for row in book["Заказы"].iter_rows(values_only=True)]
    header = next(row for row in detail if row and row[0] == "Заказ")
    assert "Сумма" in header
    assert "Площадь фасадов, м²" in header
    # строка = заказ: под шапкой есть хотя бы одна сделка
    assert any(row and row[0] and row[0] not in {"Заказ", "Заказы этих клиентов"} for row in detail)


def test_clients_xlsx_carries_the_inactive_threshold(client, auth_headers):
    """Порог «давно не покупал» уходит в файл и печатается в шапке.

    Кнопка «Клиенты, Excel» его не передавала: на экране был список на обзвон,
    а в файле — вся база.
    """
    _priced_order(
        client,
        auth_headers,
        email="stale@example.com",
        article="ОТЧ-С",
        price="900.00",
        width=1000,
        height=1000,
        quantity=1,
    )

    response = client.get(
        "/api/v1/reports/clients.xlsx?period=all&inactive_days=1", headers=auth_headers
    )

    assert response.status_code == 200
    _, _, rows = _sheet_values(response.content)
    assert _pairs(rows)["Не заказывал дольше, дней"] == 1


def test_clients_xlsx_totals_match_the_filtered_rows(client, auth_headers):
    """Шапка и таблица одного файла считаются по одному и тому же отбору."""
    _priced_order(
        client,
        auth_headers,
        email="match@example.com",
        article="ОТЧ-М",
        price="900.00",
        width=1000,
        height=1000,
        quantity=1,
    )

    # порог заведомо больше возраста заказа — под фильтр не попадает никто
    response = client.get(
        "/api/v1/reports/clients.xlsx?period=all&inactive_days=3650", headers=auth_headers
    )

    _, _, rows = _sheet_values(response.content)
    labels = _pairs(rows)
    assert labels["Клиентов"] == 0
    assert labels["Сумма"] == 0.0


# --- заказы построчно --------------------------------------------------------


def test_orders_xlsx_row_per_order(client, auth_headers):
    """Строка = заказ: номер, клиент, сумма и площадь."""
    _priced_order(
        client,
        auth_headers,
        email="row@example.com",
        article="ОТЧ-З",
        price="1500.00",
        width=1000,
        height=1000,
        quantity=1,
    )

    on_screen = client.get("/api/v1/reports/orders?period=all", headers=auth_headers).json()["data"]
    response = client.get("/api/v1/reports/orders.xlsx?period=all", headers=auth_headers)

    assert response.status_code == 200
    book, sheet, rows = _sheet_values(response.content)
    assert book.sheetnames == ["Заказы"]

    header = next(row for row in rows if row and row[0] == "Заказ")
    numbers = [row[0] for row in rows if row and row[0] and row[0] not in {"Заказ", "Итого"}]
    assert on_screen["rows"][0]["order_number"] in numbers
    assert "Площадь материала, м²" in header

    total_row = next(row for row in rows if row and row[0] == "Итого")
    assert total_row[header.index("Сумма")] == float(on_screen["total"]["amount"])


def test_orders_xlsx_is_empty_for_a_quiet_period(client, auth_headers):
    """Пустой период — файл открывается, в нём только шапка."""
    response = client.get(
        "/api/v1/reports/orders.xlsx?date_from=2000-01-01&date_to=2000-01-02",
        headers=auth_headers,
    )

    assert response.status_code == 200
    _, _, rows = _sheet_values(response.content)
    assert _pairs(rows)["Заказов"] == 0
    assert not any(row and row[0] == "Итого" for row in rows)


# --- форматы ячеек -----------------------------------------------------------


def test_counters_in_the_header_are_not_money(client, auth_headers):
    """«Заказов: 3», а не «3,00».

    Денежный формат вешался на всё числовое подряд, включая счётчики и площади:
    заказчик читала «3,00» в шапке как сумму в рублях.
    """
    _priced_order(
        client,
        auth_headers,
        email="fmt@example.com",
        article="ОТЧ-Ф",
        price="900.00",
        width=1000,
        height=1000,
        quantity=1,
    )

    response = client.get("/api/v1/reports/volume.xlsx?period=all", headers=auth_headers)
    book = load_workbook(io.BytesIO(response.content))
    sheet = book["Объёмы"]

    cells = {row[0].value: row[1] for row in sheet.iter_rows(min_col=1, max_col=2) if row[0].value}
    assert cells["Заказов"].number_format != "# ##0.00"
    assert cells["Заказов"].value == 1
    # Площадь — со своим форматом, не денежным: это метры, а не рубли
    assert cells["Площадь фасадов, м²"].number_format == "0.00"
    # Деньги денежным форматом остаются
    assert cells["Сумма"].number_format == "# ##0.00"
