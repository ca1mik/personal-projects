from __future__ import annotations

import hashlib
import io
import json
import os
import re
import secrets
import sqlite3
import uuid
import zipfile
from datetime import datetime
from pathlib import Path
from urllib.parse import urlsplit

from fastapi import APIRouter, Depends, HTTPException, Request, Response
from fastapi.responses import FileResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.main import app
from app.core.security import hash_password
from app.db.session import get_db, engine
from app.modules.auth.dependencies import get_current_user
from app.modules.auth.service import AuthService
from app.modules.users.models import User
from app.modules.orders.models import Order
from app.modules.orders.service import OrderService
from app.modules.orders.schemas import OrderCreate, OrderPositionsReplace, OrderPositionWrite
from app.modules.orders.positions_service import OrderPositionService
from app.modules.clients.models import Client

from .store import store, now
from .crm import Bitrix, DemoBitrix, CrmError, validate_webhook, scope_for, synchronize

router=APIRouter(prefix='/connect/api')
STATIC=Path(__file__).parent/'static'
DEMO=os.environ.get('FASAD_CONNECT_DEMO')=='1'


@app.middleware('http')
async def protect_browser_writes(request: Request, call_next):
    if request.url.path.startswith('/connect/api') and request.method in ('POST','PUT','PATCH','DELETE'):
        origin=request.headers.get('origin')
        if origin and urlsplit(origin).netloc != request.headers.get('host'):
            from fastapi.responses import JSONResponse
            return JSONResponse({'detail':'Обновите страницу приложения и повторите действие'},status_code=403)
    response=await call_next(request)
    if request.url.path.startswith('/connect'):
        response.headers['Cache-Control']='no-store'
        response.headers['X-Content-Type-Options']='nosniff'
        response.headers['Referrer-Policy']='same-origin'
        response.headers['Content-Security-Policy']="default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'self'; base-uri 'none'; form-action 'self'"
    return response


def actor(request:Request, db:Session=Depends(get_db)):
    token=request.headers.get('authorization','').removeprefix('Bearer ') or request.cookies.get('access_token')
    if not token: raise HTTPException(401,'Войдите в рабочее пространство')
    user=get_current_user(token,db)
    if user.role.code!='admin' and not any(p.code=='orders.manage' for p in user.role.permissions):
        raise HTTPException(403,'Для этого раздела нужно право вести заказы')
    return user


def admin(user=Depends(actor)):
    if user.role.code!='admin': raise HTTPException(403,'Настройки доступны администратору')
    return user


def set_cookie(response, token):
    response.set_cookie('access_token',token,httponly=True,samesite='strict',max_age=28800,path='/')


class Setup(BaseModel):
    company:str=Field(min_length=2,max_length=120)
    login:str=Field(min_length=3,max_length=64,pattern=r'^[A-Za-z0-9_.-]+$')
    password:str=Field(min_length=12,max_length=120)


class Login(BaseModel):
    login:str=Field(max_length=100)
    password:str=Field(max_length=120)


def config_public():
    s=store(); c=s.config()
    return {**c,'webhook_configured':bool(s.get('webhook')),'portal':urlsplit(s.webhook()).hostname if s.get('webhook') else None}


@router.get('/session')
def session(request:Request,db:Session=Depends(get_db)):
    user=None
    try:
        u=actor(request,db)
        user={'name':u.full_name,'login':u.login,'admin':u.role.code=='admin'}
    except Exception: pass
    return {'user':user,'demo':DEMO,'needs_setup':not store().get('owner_ready',False),'company':store().config()['company'],'version':'0.1.0'}


@router.post('/setup')
def setup(payload:Setup,response:Response,db:Session=Depends(get_db)):
    if DEMO: raise HTTPException(409,'В демонстрации настройка владельца не требуется')
    with store().lock:
        if store().get('owner_ready'): raise HTTPException(409,'Рабочее пространство уже настроено')
        users=db.scalars(select(User).options(selectinload(User.role))).all()
        owner=next(x for x in users if x.role.code=='admin')
        owner.login=payload.login; owner.full_name=payload.company; owner.password_hash=hash_password(payload.password)
        for u in users: u.is_active=u.id==owner.id
        db.commit()
        c=store().config(); c['company']=payload.company; store().set('config',c); store().set('owner_ready',True)
        result=AuthService.login(db,payload.login,payload.password)
        set_cookie(response,result.token.access_token)
        store().event(None,'info','Рабочее пространство настроено')
    return {'ok':True}


@router.post('/login')
def login(payload:Login,response:Response,db:Session=Depends(get_db)):
    result=AuthService.login(db,payload.login,payload.password)
    set_cookie(response,result.token.access_token)
    return {'ok':True}


@router.post('/demo-login')
def demo_login(response:Response,db:Session=Depends(get_db)):
    if not DEMO: raise HTTPException(404,'Не найдено')
    from app.core.security import create_access_token
    owner=db.scalar(select(User).where(User.login=='owner'))
    token=create_access_token(str(owner.id))
    set_cookie(response,token)
    return {'ok':True}


@router.post('/logout')
def logout(response:Response):
    response.delete_cookie('access_token',path='/')
    return {'ok':True}


@router.get('/settings')
def settings(_=Depends(admin)): return config_public()


@router.get('/workspace')
def workspace(_=Depends(actor)):
    c=config_public()
    return {key:c.get(key) for key in ('company','mode','portal')}


class SettingsInput(BaseModel):
    company:str=Field(min_length=2,max_length=120)
    mode:str=Field(pattern='^(demo|bitrix)$')
    webhook:str=Field(default='',max_length=500)
    public_url:str=Field(default='',max_length=300)
    status_field:str=Field(default='',max_length=60)
    stage_map:dict[str,str]=Field(default_factory=dict)


@router.post('/settings')
def save_settings(p:SettingsInput,_=Depends(admin)):
    if DEMO and p.mode!='demo': raise HTTPException(409,'Реальная CRM подключается в рабочем режиме через start.bat')
    if p.webhook:
        try: webhook=validate_webhook(p.webhook)
        except (ValueError,TypeError) as e: raise HTTPException(422,str(e)) from None
    else: webhook=''
    if p.mode=='bitrix' and not webhook and not store().get('webhook'):
        raise HTTPException(422,'Укажите входящий вебхук')
    if p.public_url:
        u=urlsplit(p.public_url)
        if u.scheme!='https' or not u.hostname or u.query or u.fragment or u.username or u.password or u.path not in ('','/'):
            raise HTTPException(422,'Укажите HTTPS адрес сервера без пути и параметров')
    if p.status_field and not re.fullmatch('UF_CRM_[A-Z0-9_]+',p.status_field): raise HTTPException(422,'Код поля должен начинаться с UF_CRM_')
    if len(p.stage_map)>100 or any(not re.fullmatch('[A-Za-z0-9_:-]{1,80}',v) for v in p.stage_map.values()):
        raise HTTPException(422,'Проверьте коды стадий')
    if webhook: store().save_webhook(webhook)
    store().set('config',p.model_dump(exclude={'webhook'}))
    store().event(None,'info','Настройки подключения обновлены')
    return config_public()


@router.post('/test-connection')
def test_connection(_=Depends(admin)):
    c=store().config()
    try:
        crm=DemoBitrix(store()) if c['mode']=='demo' else Bitrix(store().webhook())
        fields=crm.call('crm.deal.fields')['result']
        missing=[x for x in ('OPPORTUNITY','ORIGIN_ID','ORIGINATOR_ID') if x not in fields]
        if c.get('status_field') and c['status_field'] not in fields: missing.append(c['status_field'])
        if missing: raise CrmError('В CRM не найдены поля: '+', '.join(missing))
    except (CrmError,ValueError) as e: raise HTTPException(422,str(e)) from None
    store().event(None,'success','Проверка связи выполнена')
    return {'ok':True,'message':'Тестовая CRM готова' if c['mode']=='demo' else 'Битрикс24 отвечает. Права чтения полей проверены.'}


def snapshot(order):
    positions=[{'name':p.name,'article':p.article or '','width':str(p.width or ''),'height':str(p.height or ''),'quantity':str(p.quantity),'thickness':str(p.thickness or ''),'color':p.color or '','milling':p.milling or '','area':str(p.area_sqm),'price':str(p.unit_price),'amount':str(p.amount)} for p in order.positions]
    return {'id':str(order.id),'number':order.order_number,'client':order.client.organization or order.client.full_name if order.client else 'Без клиента',
        'amount':str(order.final_amount),'status':order.status.name if order.status else 'Новый','status_code':order.status.code if order.status else 'new',
        'direction':order.direction.name if order.direction else 'Без направления','due_date':str(order.due_date or ''),'positions':positions}


def quote_xlsx(data):
    wb=Workbook();ws=wb.active;ws.title='Расчёт заказа'
    ws.append([store().config()['company']]);ws.append([f"Расчёт заказа № {data['number']}"])
    ws.append(['Клиент',data['client']]);ws.append(['Направление',data['direction']]);ws.append(['Статус',data['status']]);ws.append([])
    ws.append(['№','Наименование','Длина мм','Ширина мм','Кол-во','Толщина мм','Цвет','Площадь м²','Цена ₽','Сумма ₽'])
    for i,row in enumerate(data['positions'],1):
        def number(k): return float(row[k]) if row[k] else None
        ws.append([i,row['name'],number('height'),number('width'),number('quantity'),number('thickness'),row['color'],number('area'),number('price'),number('amount')])
    ws.append([]);ws.append(['Итого по заказу',float(data['amount'])]);ws.append(['Валюта','RUB'])
    for row in ws:
        for cell in row:
            if isinstance(cell.value,str): cell.data_type='s'
            cell.font=Font(name='Arial',size=10)
            cell.alignment=Alignment(vertical='center',wrap_text=True)
    for cell in ws[7]: cell.fill=PatternFill('solid',fgColor='26384C');cell.font=Font(name='Arial',bold=True,color='FFFFFF',size=10)
    for col,w in zip('ABCDEFGHIJ',[6,32,14,14,10,14,23,15,16,18]):ws.column_dimensions[col].width=w
    ws.freeze_panes='A8';ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_title_rows='1:7'
    out=io.BytesIO();wb.save(out);return out.getvalue()


@router.get('/orders')
def orders(db:Session=Depends(get_db),_=Depends(actor)):
    results=db.scalars(select(Order).where(Order.deleted_at.is_(None)).order_by(Order.created_at.desc()).limit(500)).all()
    c=store().config();scope=scope_for(c,store().webhook() if c['mode']=='bitrix' else '')
    rows=[]
    for order in results:
        item=snapshot(order);item['sync']=store().link(scope,order.id);item['positions_count']=len(item['positions'])
        item['area']=str(sum(float(p['area']) for p in item['positions']));rows.append(item)
    return {'items':rows,'mode':c['mode']}


@router.get('/orders/{order_id}/quote.xlsx')
def quote(order_id:uuid.UUID,db:Session=Depends(get_db),_=Depends(actor)):
    data=snapshot(OrderService.get_order(db,order_id))
    return Response(quote_xlsx(data),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="Fasad-{order_id}.xlsx"'})


class SyncInput(BaseModel):
    existing_deal_id:int|None=Field(default=None,gt=0)


@router.post('/orders/{order_id}/sync')
def sync_order(order_id:uuid.UUID,p:SyncInput,db:Session=Depends(get_db),_=Depends(actor)):
    data=snapshot(OrderService.get_order(db,order_id))
    if not data['positions']: raise HTTPException(422,'Сначала добавьте позиции расчёта в заказ')
    filename='Fasad-'+re.sub(r'[^A-Za-z0-9А-Яа-яЁё_-]','-',data['number'])+'.xlsx'
    try: return synchronize(store(),data,quote_xlsx(data),filename,p.existing_deal_id)
    except CrmError as e: raise HTTPException(422,str(e)) from None


@router.get('/journal')
def journal(_=Depends(actor)):
    with store().db() as db: return {'items':[dict(x) for x in db.execute('SELECT * FROM events ORDER BY id DESC LIMIT 200')]}


@router.get('/demo-deals')
def demo_deals(_=Depends(actor)):
    with store().db() as db:
        deals=[]
        for d in db.execute('SELECT * FROM demo_deals ORDER BY id DESC'):
            comments=[{'id':r['id'],**json.loads(r['fields'])} for r in db.execute('SELECT * FROM demo_comments WHERE deal_id=?',(d['id'],))]
            # Expose filenames and size, never base64 payloads, in the interface.
            for c in comments: c['FILES']=[{'name':f[0],'bytes':len(f[1])*3//4} for f in c.get('FILES',[])]
            deals.append({'id':d['id'],**json.loads(d['fields']),'comments':comments})
        return {'items':deals}


@router.get('/backup')
def backup(_=Depends(admin)):
    import tempfile
    with tempfile.TemporaryDirectory() as temp:
        with store().lock:
            for source,name in [(Path(engine.url.database),'fasad.sqlite'),(store().path,'connect.sqlite')]:
                s=sqlite3.connect(source);d=sqlite3.connect(Path(temp)/name)
                try:s.backup(d)
                finally:d.close();s.close()
            out=io.BytesIO()
            with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
                for f in Path(temp).iterdir():z.write(f,f.name)
                z.writestr('connect.key', (store().directory/'connect.key').read_bytes())
                configuration=store().directory/'config.json'
                if configuration.exists():z.write(configuration,'config.json')
                files=store().directory/'files'
                if files.exists():
                    for f in files.rglob('*'):
                        if f.is_file() and f.resolve().is_relative_to(files.resolve()):z.write(f,('files'/f.relative_to(files)).as_posix())
                z.writestr('README.txt','Резервная копия Фасад Connect 0.1. Остановите приложение. Сохраните текущую папку runtime. Восстанавливайте fasad.sqlite, connect.sqlite, connect.key, config.json и files в папку runtime/demo или runtime/pilot выбранной установки. Не смешивайте demo и pilot. Для проверки сначала используйте отдельную копию продукта. Архив содержит ключ подключения: храните его как рабочие данные предприятия.')
    return Response(out.getvalue(),media_type='application/zip',headers={'Content-Disposition':'attachment; filename="Fasad-backup.zip"'})


app.include_router(router)


@app.get('/connect')
def index():return FileResponse(STATIC/'index.html')


@app.get('/connect/assets/{name}')
def asset(name:str):
    if name not in ('app.js','style.css'):raise HTTPException(404)
    return FileResponse(STATIC/name)
