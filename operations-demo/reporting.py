"""Portable application XLSX export and cross-tab calculations from verified records.

The server must work on a reviewer's PC without a Codex artifact runtime.
Openpyxl is an application dependency, not a dependency on the original Excel files.
"""
from collections import defaultdict
from datetime import date
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.chart import BarChart,Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.worksheet.table import Table,TableStyleInfo
from domain import AREAS,quantity

def pivot(report,area,group_by):
    allowed={'date','batch_id'} if area=='welding' else {'date','employee','batch_id'} if area=='production' else {'date','employee'}
    if area not in AREAS or group_by not in allowed:raise ValueError('Такая группировка недоступна для выбранного участка')
    groups=defaultdict(list)
    for r in report['records']+report['undated']:
        if r['area']==area:groups[str(r['values'].get(group_by) or 'Не заполнено')].append(r)
    rows=[]
    for label,items in sorted(groups.items()):
        good=[r for r in items if r['valid']];suffix='qty' if area=='production' else 'docs'
        plan=sum(quantity(r['values']['plan_'+suffix]) for r in good) if area!='welding' else None
        fact=sum(quantity(r['values']['fact_'+suffix]) for r in good) if area!='welding' else sum(quantity(r['values']['need_qty']) for r in good)
        complete=len(items)==len(good)
        rows.append({'group':label,'plan':plan,'fact':fact,'ratio':fact/plan if complete and plan else None,
                     'rows':len(items),'valid_rows':len(good),'complete':complete,'record_ids':[r['id'] for r in items]})
    return {'area':area,'group_by':group_by,'rows':rows}

def build_xlsx(report,area,group_by):
    wb=Workbook();ws=wb.active;ws.title='Обзор'
    dark='163E35';teal='2C8067';pale='E8F2EC';amber='FFF2DA'
    def cell(sheet,row,col,value,fmt=None):
        c=sheet.cell(row,col,value)
        # Imported text must remain text even when it begins with '='.
        if isinstance(value,str):c.data_type='s'
        if fmt:c.number_format=fmt
        return c
    def line(sheet,row,values):
        for col,value in enumerate(values,1):cell(sheet,row,col,value)
    def title(sheet,text):
        sheet.sheet_view.showGridLines=False
        c=cell(sheet,2,1,text);c.font=Font(name='Arial',size=16,bold=True,color=dark)
    def table(sheet,header,rows,start=5,name='Data'):
        line(sheet,start,header)
        for i,row in enumerate(rows,start+1):line(sheet,i,row)
        end=start+len(rows)
        for c in sheet[start]:
            if c.column<=len(header):c.fill=PatternFill('solid',fgColor=dark);c.font=Font(name='Arial',size=11,bold=True,color='FFFFFF');c.alignment=Alignment(wrap_text=True,vertical='center')
        sheet.row_dimensions[start].height=33
        if rows:
            t=Table(displayName=name,ref=f'A{start}:{sheet.cell(end,len(header)).column_letter}{end}')
            t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium4',showRowStripes=True);sheet.add_table(t)
        sheet.freeze_panes=f'C{start+1}'
        return end
    title(ws,'Производство по участкам')
    cell(ws,3,1,'Период');cell(ws,3,2,date.fromisoformat(report['start']),'dd.mm.yyyy');cell(ws,3,3,date.fromisoformat(report['end']),'dd.mm.yyyy')
    overview=[]
    for a,t in report['totals'].items():overview.append([AREAS[a],'док.' if a=='docs' else 'ед.',t.get('known_plan'),t.get('known_fact',t.get('known_need')),t.get('ratio'),t['valid_rows'],t['rows'],'Полные' if t['complete'] else 'Частичные'])
    end=table(ws,['Участок','Единица','План','Факт / потребность','Выполнение','Проверено строк','Всего строк','Данные'],overview,name='Overview')
    for row in range(6,end+1):ws.cell(row,5).number_format='0.0%'
    notes=['Суммы относятся только к проверенным строкам; частичная сумма не является итогом всего участка.',
           'Выполнение = общий факт / общий план. При ошибках общий процент не показывается.',
           'Сварка связана с заготовкой по партии внутри периода. Документация независима.',
           'Записи с неверной датой включены в первичные данные и сводную с явными замечаниями.',
           'Сводная на отдельном листе — рассчитанная группировка, не объект Excel PivotTable. Для анализа исходных данных доступны фильтры.',
           'Диаграммы построены по проверенным дневным значениям; пропуски не считаются нулями.']
    for i,text in enumerate(notes,end+3):cell(ws,i,1,text)
    summary=wb.create_sheet('Сводная');title(summary,'Сводка · '+AREAS[area])
    names={'employee':'Сотрудник','date':'Дата','batch_id':'Партия'};cell(summary,3,1,'Группировка: '+names[group_by])
    cross=pivot(report,area,group_by)
    def group_value(value):
        if group_by=='date':
            try:return date.fromisoformat(value)
            except ValueError:pass
        return value
    crossrows=[[group_value(r['group']),r['plan'],r['fact'],r['ratio'],r['valid_rows'],r['rows'],'Полные' if r['complete'] else 'Частичные'] for r in cross['rows']]
    last=table(summary,[names[group_by],'План','Факт / потребность','Выполнение','Проверено','Всего','Данные'],crossrows,name='Grouped')
    for row in range(6,last+1):
        summary.cell(row,4).number_format='0.0%'
        if group_by=='date' and isinstance(summary.cell(row,1).value,date):summary.cell(row,1).number_format='dd.mm.yyyy'
    batches=wb.create_sheet('Партии');title(batches,'Обеспеченность сварки')
    table(batches,['Партия','Заготовлено','Нужно','Дефицит (+) / избыток (−)','Состояние'],[[b['batch_id'],b['fact'],b['need'],b['gap'],'; '.join(b['issues']) or 'Достоверная сверка'] for b in report['batches']],name='Batches')
    raw=wb.create_sheet('Первичные');title(raw,'Первичные записи и качество данных')
    rawrows=[]
    def number(value):
        if value in (None,''):return None
        try:return float(value)
        except (ValueError,TypeError):return str(value)
    for r in report['records']+report['undated']:
        d=r['values'];v=d.get('date')
        try:v=date.fromisoformat(v)
        except (ValueError,TypeError):pass
        rawrows.append([AREAS[r['area']],d.get('record_id'),v,d.get('batch_id'),d.get('employee'),number(d.get('plan_qty',d.get('plan_docs'))),number(d.get('fact_qty',d.get('fact_docs'))),number(d.get('need_qty')),'Да' if r['valid'] else 'Нет','; '.join(r['issues'])])
    table(raw,['Участок','Номер записи','Дата','Партия','Сотрудник','План','Факт','Потребность','Проверено','Замечания'],rawrows,name='SourceRecords')
    for i,r in enumerate(rawrows,6):
        raw.cell(i,3).number_format='dd.mm.yyyy'
        if r[8]=='Нет':
            for c in raw[i]:c.fill=PatternFill('solid',fgColor=amber)
    daily=wb.create_sheet('Диаграммы');title(daily,'План и факт по дням')
    for idx,a in enumerate(x for x in ('production','docs') if x in report['totals'] and area!='welding'):
        start=5+idx*(max(len(report['daily']),2)+22)
        line(daily,start,[AREAS[a]+' · дата','План','Факт','День'])
        for i,p in enumerate(report['daily'],start+1):
            day=date.fromisoformat(p['date'])
            line(daily,i,[day,p[a+'_plan'],p[a],day.strftime('%d.%m')])
            daily.cell(i,1).number_format='dd.mm.yyyy'
        if report['daily']:
            chart=BarChart();chart.title=AREAS[a]+' · план и факт';chart.y_axis.title='документов' if a=='docs' else 'единиц';chart.x_axis.title='Дата';chart.style=10;chart.height=9;chart.width=23;chart.display_blanks='gap'
            chart.add_data(Reference(daily,min_col=2,max_col=3,min_row=start,max_row=start+len(report['daily'])),titles_from_data=True)
            for series,label,colour in zip(chart.series,('План','Факт'),('CCDCD4','39876A')):
                series.tx=SeriesLabel(v=label);series.graphicalProperties.solidFill=colour
            chart.legend.position='b'
            chart.set_categories(Reference(daily,min_col=4,min_row=start+1,max_row=start+len(report['daily'])))
            daily.add_chart(chart,f'E{start}')
    if area=='welding':
        # Welding has no invented plan/fact relation: chart the entered demand only.
        start=5;line(daily,start,['Партия','Потребность'])
        for i,r in enumerate(cross['rows'],6):line(daily,i,[r['group'],r['fact'] if r['complete'] else None])
        if cross['rows']:
            chart=BarChart();chart.title='Потребность сварки';chart.add_data(Reference(daily,min_col=2,min_row=5,max_row=5+len(cross['rows'])),titles_from_data=True);chart.set_categories(Reference(daily,min_col=1,min_row=6,max_row=5+len(cross['rows'])));chart.series[0].tx=SeriesLabel(v='Потребность');chart.legend.position='b';chart.width=23;chart.height=9;daily.add_chart(chart,'E5')
    for sheet in wb:
        for row in sheet:
            for c in row:
                if c.row!=2 and not str(c.fill.fgColor.rgb).endswith(dark):c.font=Font(name='Arial',size=11,color=dark)
                c.alignment=Alignment(vertical='center',wrap_text=c.column in (8,10))
                if isinstance(c.value,(int,float)):
                    c.number_format=('#,##0' if float(c.value).is_integer() else '0.00') if c.number_format=='General' else c.number_format
        for col in range(1,sheet.max_column+1):sheet.column_dimensions[sheet.cell(1,col).column_letter].width=22 if col<6 else 20
        sheet.column_dimensions['A'].width=31
    raw.column_dimensions['E'].width=27;raw.column_dimensions['J'].width=75;raw.column_dimensions['C'].width=16;batches.column_dimensions['E'].width=65
    output=BytesIO();wb.save(output);return output.getvalue()
