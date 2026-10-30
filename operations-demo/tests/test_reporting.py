from io import BytesIO
from datetime import datetime
from openpyxl import load_workbook
from test_acceptance import demo,client
from app import create_app

def test_pivot_groups_period_and_missing_data(demo):
    c,_,_=demo
    p=c.get('/api/pivot?start=2026-10-05&end=2026-10-14&area=production&group_by=employee').json()
    anna=next(r for r in p['rows'] if r['group']=='Анна Соколова')
    assert anna['rows']==5 and anna['valid_rows']==2 and anna['fact']==190 and anna['ratio'] is None
    two=c.get('/api/pivot?start=2026-10-05&end=2026-10-06&area=production&group_by=employee').json()
    assert len(two['rows'])==1 and two['rows'][0]['ratio']==.95
    assert c.get('/api/pivot?start=2026-10-05&end=2026-10-14&area=welding&group_by=employee').status_code==422

def test_xlsx_has_typed_dates_percentages_charts_and_source_errors(demo):
    c,_,_=demo
    r=c.get('/api/export.xlsx?start=2026-10-05&end=2026-10-14&area=production&group_by=employee')
    assert r.status_code==200 and 'attachment' in r.headers['content-disposition']
    wb=load_workbook(BytesIO(r.content))
    assert wb.sheetnames==['Обзор','Сводная','Партии','Первичные','Диаграммы']
    assert isinstance(wb['Обзор']['B3'].value,datetime)
    assert wb['Обзор']['E8'].value==18/26 and wb['Обзор']['E8'].number_format=='0.0%'
    assert wb['Обзор']['A5'].font.color.rgb.endswith('FFFFFF')
    assert wb['Первичные']['G9'].value==-5
    assert wb['Первичные']['C6'].number_format=='dd.mm.yyyy'
    assert len(wb['Диаграммы']._charts)==2
    assert wb['Сводная'].tables and wb['Первичные'].tables
    assert wb['Первичные'].column_dimensions['C'].width>=16

def test_xlsx_is_scoped_and_never_executes_imported_text(demo):
    c,db,_=demo
    c.post('/api/records/docs',json={'values':{'record_id':'D-test','date':'2026-99-99','employee':'=SUM(A1:A2)','plan_docs':'1','fact_docs':'1'}})
    worker=client(create_app(db),'docs')
    assert worker.get('/api/export.xlsx?start=2026-10-05&end=2026-10-14&area=production&group_by=employee').status_code==403
    r=worker.get('/api/export.xlsx?start=2026-10-05&end=2026-10-14&area=docs&group_by=date');assert r.status_code==200
    wb=load_workbook(BytesIO(r.content))
    assert all(row[0].value in ('Документация',None) for row in wb['Первичные'].iter_rows(min_row=6))
    text=next(cell for row in wb['Первичные'] for cell in row if cell.value=='=SUM(A1:A2)')
    assert text.data_type=='s'
    assert c.get('/api/export.xlsx?start=2026-10-05&end=2026-10-14&area=unknown').status_code==422

def test_role_matrix_matches_actual_permissions(demo):
    c,db,_=demo
    roles={r['id']:r for r in c.get('/api/roles').json()['roles']}
    assert roles['admin']['manage_users'] and len(roles['admin']['edit_areas'])==3
    assert roles['director']['edit_areas']==[] and len(roles['director']['read_areas'])==3
    assert roles['production']['edit_areas']==['production']
    assert client(create_app(db),'production').get('/api/roles').status_code==403
